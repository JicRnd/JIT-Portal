from __future__ import annotations
from datetime import date, time
import csv
from decimal import Decimal, InvalidOperation
from fractions import Fraction
from pathlib import Path

import openpyxl

D = Decimal


STANDARD_SERIES_MODIFICATION_COLUMNS = {
    'rod_extension': {'per_length': 'Rod_Extension_per_inch'},
    'stop_tube': {'base': 'Stop_Tub_Base_price', 'per_length': 'Stop_tube_per_inch'},
    'stainless_rod': {
        'base': 'Stainless_Steel_pisiton_Rod_base_price',
        'per_stroke': 'Stainless_Steel_pisiton_Rod_per_inch',
    },
    'chrome_bore': {'base': 'Chrome_Bore_base_price', 'per_stroke': 'Chrome_Bore_price_inch'},
    'extra_thread': {'per_length': 'extra_thread'},
}


def dec(value) -> Decimal:
    if value is None or value == "":
        return D("0")
    if isinstance(value, Decimal):
        return value
    return D(str(value).strip())


def parse_dimension(value) -> Decimal:
    """Parse numeric values and Excel-style mixed fractions such as 1 3/8\"."""
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)):
        return D(str(value))
    s = str(value).strip().replace('"', '')
    try:
        return D(s)
    except InvalidOperation:
        pass
    if ' ' in s:
        whole, frac = s.split(None, 1)
        f = Fraction(frac)
        return D(whole) + (D(f.numerator) / D(f.denominator))
    f = Fraction(s)
    return D(f.numerator) / D(f.denominator)


class PricingData:
    def __init__(self, data_dir: str | Path):
        self.data_dir = Path(data_dir)
        self.standard_series_pricing = self._read_series_pricing_xlsx('standard_series_pricing.xlsx')
        self.metric_series_pricing = self._read_series_pricing_xlsx('metric_series_pricing.xlsx')
        self.series_pricing = self.standard_series_pricing + self.metric_series_pricing
        self.mod_rates = self._read_modification_rate_matrix()
        self.special_parts = self._read('special_part_catalog.csv')
        self.ph = self._read('ph_position_sensing.csv')
        self.ph_bands = self._read('ph_sensor_bands.csv')
        self.va_base = self._read('va_base_pricing.csv')
        self.va_mods = self._read('va_modifications.csv')
        self.accessories = self._read('accessory_mountings.csv')
        self.accessory_catalog = self._read('accessory_catalog_v12.csv')
        self.accessory_selectors = self._read('acc_selector_rules.csv')
        self.rod_boot_rates = self._read('rod_boot_rates.csv')

    def _read(self, name: str):
        workbook_path = self.data_dir / f'{Path(name).stem}.xlsx'
        if workbook_path.is_file():
            return self._read_xlsx_table(workbook_path)
        with (self.data_dir / name).open(encoding='utf-8-sig', newline='') as f:
            return list(csv.DictReader(f))

    def _read_modification_rate_matrix(self):
        matrix_name = 'unused/metric_modification_rate_matrix.xlsx'
        if not (self.data_dir / matrix_name).is_file():
            matrix_name = 'metric_modification_rate_matrix.xlsx'
        workbook_names = ((matrix_name, 'mm'),)
        workbook_paths = [self.data_dir / name for name, _ in workbook_names]
        if not all(path.is_file() for path in workbook_paths):
            missing = [str(path) for path in workbook_paths if not path.is_file()]
            raise FileNotFoundError(
                'Required modification rate workbooks are missing: ' + ', '.join(missing)
            )

        rows = []
        for (workbook_name, expected_unit), workbook_path in zip(workbook_names, workbook_paths):
            workbook_rows = self._read_xlsx_table(workbook_path)
            if any(row.get('unit_system') != expected_unit for row in workbook_rows):
                raise ValueError(f'Unexpected unit_system in {workbook_path}')
            rows.extend(workbook_rows)
        return rows

    @staticmethod
    def _read_xlsx_table(workbook_path: Path):
        workbook = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)
        try:
            worksheet = workbook[workbook_path.stem] if workbook_path.stem in workbook.sheetnames else workbook.active
            header_values = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
            headers = [str(value) if value is not None else '' for value in (header_values or ())]
            rows = []
            for values in worksheet.iter_rows(min_row=2, values_only=True):
                record = {}
                for header, value in zip(headers, values):
                    if isinstance(value, date):
                        value = f'{value.month}/{value.day}-{value.year % 100:02d}'
                    elif isinstance(value, time):
                        raise ValueError(f'Unsupported time-typed value in {workbook_path}')
                    record[header] = '' if value is None else str(value)
                rows.append(record)
            return rows
        finally:
            workbook.close()

    def _read_series_pricing_xlsx_files(self):
        rows = []
        for workbook_name in ('standard_series_pricing.xlsx', 'metric_series_pricing.xlsx'):
            rows.extend(self._read_series_pricing_xlsx(workbook_name))
        return rows

    def _read_series_pricing_xlsx(self, workbook_name):
        workbook_path = self.data_dir / workbook_name
        if not workbook_path.is_file():
            unit = 'mm' if workbook_name.startswith('metric_') else 'inch'
            raise FileNotFoundError(
                f'Required {unit} series pricing workbook is missing: {workbook_path}. '
                f'{workbook_name} is required for {unit} calculator pricing; '
                'the legacy series row-rate workbook is not a fallback.'
            )
        workbook = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)
        try:
            if 'series_pricing' not in workbook.sheetnames:
                raise ValueError(f'Missing series_pricing worksheet in {workbook_path}')
            worksheet = workbook['series_pricing']
            header_values = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
            headers = [str(value) if value is not None else '' for value in (header_values or ())]
            if 'base_price' not in headers or 'base_price_formula' not in headers:
                raise ValueError(f'Unexpected series_pricing headers in {workbook_path}')

            rows = []
            for row_number, values in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
                record = {}
                for header, value in zip(headers, values):
                    if header == 'base_price_formula':
                        continue
                    if isinstance(value, date):
                        if header not in {'standard_thread', 'oversize_thread'}:
                            raise ValueError(
                                f'Unsupported date-typed value in {workbook_path} '
                                f'at row {row_number}, column {header}'
                            )
                        value = f'{value.month}/{value.day}-{value.year % 100:02d}'
                    elif isinstance(value, time):
                        raise ValueError(
                            f'Unsupported time-typed value in {workbook_path} '
                            f'at row {row_number}, column {header}'
                        )
                    record[header] = '' if value is None else str(value)
                if 'stroke_unit' not in record and 'unit_of_measure' in record:
                    record['stroke_unit'] = record['unit_of_measure']
                record['base_price_formula'] = ''
                rows.append(record)
            return rows
        finally:
            workbook.close()

    @staticmethod
    def _aliases(row):
        return {x.strip().upper() for x in row.get('series_aliases', '').split('|') if x.strip()}

    def find_series_price(self, series, bore, rod, mount):
        series = series.upper(); mount = mount.upper()
        pricing_rows = self.metric_series_pricing if series in {'IH', 'IHM', 'IMH'} else self.standard_series_pricing
        candidates = []
        for r in pricing_rows:
            if series not in self._aliases(r):
                continue
            if parse_dimension(r['bore']) != bore or parse_dimension(r['rod_diameter']) != rod:
                continue
            codes = {x.strip().upper() for x in r['mount_codes'].split('|') if x.strip()}
            if mount in codes:
                candidates.append(r)
        if len(candidates) != 1:
            raise LookupError(f'Expected one base-price row for {series}/{bore}/{rod}/{mount}, found {len(candidates)}')
        return candidates[0]

    def find_series_row(self, series, bore, rod):
        series = series.upper()
        metric = series in {'IH', 'IHM', 'IMH'}
        pricing_rows = self.metric_series_pricing if metric else self.standard_series_pricing
        expected_unit = 'mm' if metric else 'inch'
        rows = [r for r in pricing_rows
                if series in self._aliases(r)
                and parse_dimension(r['bore']) == bore
                and parse_dimension(r['rod_diameter']) == rod
                and r.get('stroke_unit') == expected_unit]
        unique_rates = {(r['stroke_rate'], r['cushion_per_end']): r for r in rows}
        if len(unique_rates) != 1:
            raise LookupError(f'Expected one row-rate record for {series}/{bore}/{rod}, found {len(unique_rates)}')
        return next(iter(unique_rates.values()))

    def find_mod_rates(self, series, bore, rod, rule_id):
        series = series.upper()
        if series not in {'IH', 'IHM', 'IMH'}:
            if rule_id not in STANDARD_SERIES_MODIFICATION_COLUMNS:
                raise LookupError(
                    f'No standard-series pricing column for modification rule {rule_id!r}'
                )
            return self._find_standard_series_mod_rates(series, bore, rod, rule_id)

        out = {}
        for r in self.mod_rates:
            aliases = {x.strip().upper() for x in r['series_group'].split('|') if x.strip()}
            if series not in aliases or r['rule_id'] != rule_id:
                continue
            if parse_dimension(r['bore']) != bore or parse_dimension(r['rod_diameter']) != rod:
                continue
            out[r['rate_name']] = dec(r['rate_value'])
        return out

    def _find_standard_series_mod_rates(self, series, bore, rod, rule_id):
        rows = [
            row for row in self.standard_series_pricing
            if series in self._aliases(row)
            and parse_dimension(row['bore']) == bore
            and parse_dimension(row['rod_diameter']) == rod
        ]
        if not rows:
            raise LookupError(
                f'No standard-series modification row for {series}/{bore}/{rod}/{rule_id}'
            )

        out = {}
        for rate_name, column in STANDARD_SERIES_MODIFICATION_COLUMNS[rule_id].items():
            values = {row.get(column, '') for row in rows}
            if len(values) != 1 or not next(iter(values), ''):
                raise LookupError(
                    f'Expected one {column} value for {series}/{bore}/{rod}, found {len(values)}'
                )
            out[rate_name] = dec(next(iter(values)))
        return out

    def find_special_part(self, part_number):
        key = part_number.strip().upper()
        rows = [r for r in self.special_parts if r['part_number'].strip().upper() == key]
        if len(rows) != 1:
            raise LookupError(f'Special part {part_number!r}: found {len(rows)} rows')
        return rows[0]

    def find_va_base(self, bore, rod):
        rows = [r for r in self.va_base if parse_dimension(r['bore']) == bore and parse_dimension(r['rod_diameter']) == rod]
        if len(rows) != 1:
            raise LookupError(f'VA base {bore}/{rod}: found {len(rows)} rows')
        return rows[0]

    def find_va_mod(self, bore, rod):
        rows = [r for r in self.va_mods if parse_dimension(r['bore']) == bore and parse_dimension(r['rod_diameter']) == rod]
        if len(rows) != 1:
            raise LookupError(f'VA mods {bore}/{rod}: found {len(rows)} rows')
        return rows[0]

    def find_ph(self, metric, bore):
        system = 'metric' if metric else 'imperial'
        rows = [r for r in self.ph if r['system'] == system and parse_dimension(r['bore']) == bore]
        if len(rows) != 1:
            raise LookupError(f'PH {system}/{bore}: found {len(rows)} rows')
        return rows[0]

    def accessory_price_by_thread(self, thread, part_type):
        rows = [r for r in self.accessories if r['thread_size'].strip() == thread and r['part_type'] == part_type]
        if len(rows) != 1:
            raise LookupError(f'Accessory {part_type}/{thread}: found {len(rows)} rows')
        return dec(rows[0]['price'])

    def find_accessory(self, lookup_type, lookup_key, part_type):
        key = str(lookup_key).strip()
        rows = [r for r in self.accessory_catalog
                if r['lookup_type'] == lookup_type
                and r['lookup_key'].strip() == key
                and r['part_type'] == part_type]
        if len(rows) > 1:
            raise LookupError(f'Accessory {part_type}/{lookup_type}/{lookup_key}: found {len(rows)} rows')
        return rows[0] if rows else None

    def find_accessory_bucket(self, series, bore, mount):
        series = str(series).strip().upper()
        mount = str(mount).strip().upper()
        matches = []
        for row in self.accessory_selectors:
            if row['series'].strip().upper() != series:
                continue
            if mount not in {item.strip().upper() for item in row['mounts'].split('|') if item.strip()}:
                continue
            lower = parse_dimension(row['bore_min'])
            upper = parse_dimension(row['bore_max'])
            lower_ok = bore >= lower if row['min_inclusive'].lower() == 'true' else bore > lower
            upper_ok = bore <= upper if row['max_inclusive'].lower() == 'true' else bore < upper
            if lower_ok and upper_ok:
                matches.append(row['bucket'])
        if len(matches) > 1:
            raise LookupError(f'Expected one accessory bucket for {series}/{bore}/{mount}, found {len(matches)}')
        return matches[0] if matches else None

    def find_rod_boot(self, rod):
        rows = [r for r in self.rod_boot_rates if parse_dimension(r['rod_diameter']) == rod]
        if len(rows) > 1:
            raise LookupError(f'Rod boot {rod}: found {len(rows)} rows')
        return rows[0] if rows else None
