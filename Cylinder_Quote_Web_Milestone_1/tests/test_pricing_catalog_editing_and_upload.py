from __future__ import annotations

import io
from decimal import Decimal

from openpyxl import Workbook

from app.db import get_session
from app.models_db import CatalogPart, InventoryPart, PriceChangeLog
from tests.test_pricing_catalog_bulk import _admin_client


def _workbook(rows, headers=("Part Number", "Sell Price")):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(list(headers))
    for row in rows:
        sheet.append(list(row))
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


def test_catalog_columns_renamed_and_vendor_column_present(tmp_path, monkeypatch):
    _, client, _ = _admin_client(tmp_path, monkeypatch)
    page = client.get("/parts-catalog").get_data(as_text=True)

    assert "<th>Price Location</th>" in page
    assert "<th>Inventory Source</th>" in page
    assert "<th>Price Source</th>" not in page
    assert "<th>Inventory Location</th>" not in page
    assert "<th>Vendors</th>" in page
    assert "Source Locations" not in page
    assert "Sell Price Source" not in page
    assert "edit-part-modal" not in page
    assert "edit-part-btn" not in page
    assert "cell-inline-edit" in page
    assert 'data-field="part_number"' in page
    assert 'data-field="description"' in page
    assert 'data-field="unit_cost"' in page
    assert 'data-field="sell_price"' in page
    assert '<th class="inventory-col">Inventory</th>' in page
    assert 'Update Inventory Bulk' in page
    assert '>Select</button>' in page
    assert 'data-field="inventory"' in page
    assert 'data-field="vendors"' in page
    assert 'data-field="source_locations"' in page


def test_edit_part_updates_every_field_and_logs_history(tmp_path, monkeypatch):
    _, client, ids = _admin_client(tmp_path, monkeypatch)
    part_id = ids["S-100"]

    loaded = client.get(f"/parts-catalog/parts/{part_id}").get_json()
    assert loaded["ok"] and loaded["part"]["part_number"] == "S-100"

    response = client.post(f"/parts-catalog/parts/{part_id}", json={
        "part_number": "S-100A",
        "description": "Viton Rod Seal 1.000 Updated",
        "category": "Seals",
        "unit_cost": "41.50",
        "sell_price": "111.00",
        "vendors": "Parker Hannifin; Motion Industries",
        "source_locations": "Sheet1!A5",
        "active": True,
    })
    data = response.get_json()
    assert response.status_code == 200 and data["ok"]

    with get_session() as session:
        row = session.get(CatalogPart, part_id)
        assert row.part_number == "S-100A"
        assert row.description == "Viton Rod Seal 1.000 Updated"
        assert row.vendors == "Parker Hannifin\nMotion Industries"
        assert row.source_locations == "Sheet1!A5"
        assert row.sell_price == Decimal("111.0000")
        assert row.unit_cost == Decimal("41.5000")
        logged = {
            entry.field_name
            for entry in session.query(PriceChangeLog).filter_by(change_type="edit").all()
        }
    assert {"part_number", "description", "vendors", "source_locations", "unit_cost", "sell_price"} <= logged


def test_inventory_inline_and_bulk_updates(tmp_path, monkeypatch):
    _, client, ids = _admin_client(tmp_path, monkeypatch)

    inline = client.post(f"/parts-catalog/parts/{ids['S-100']}", json={"inventory": "7"})
    assert inline.status_code == 200
    assert inline.get_json()["inventory"] == 7

    bulk = client.post("/parts-catalog/inventory/bulk", json={
        "ids": [ids["S-100"], ids["S-200"]],
        "inventory": "12",
    })
    assert bulk.status_code == 200
    assert bulk.get_json()["count"] == 2

    with get_session() as session:
        assert session.get(CatalogPart, ids["S-100"]).inventory is None
        assert session.query(InventoryPart).filter_by(part_number="S-100").one().inventory == 12
        assert session.query(InventoryPart).filter_by(part_number="S-200").one().inventory == 12


def test_edit_part_rejects_duplicate_part_number(tmp_path, monkeypatch):
    _, client, ids = _admin_client(tmp_path, monkeypatch)
    response = client.post(f"/parts-catalog/parts/{ids['S-100']}", json={"part_number": "S-200"})
    assert response.status_code == 400
    assert "already exists" in response.get_json()["error"]


def test_workbook_upload_only_changes_different_prices(tmp_path, monkeypatch):
    _, client, ids = _admin_client(tmp_path, monkeypatch)
    upload = _workbook([
        ("S-100", 125.00),   # different -> update
        ("S-200", 57.50),    # same -> untouched
        ("NOPE-1", 9.99),    # unknown -> skipped
    ])

    preview = client.post(
        "/parts-catalog/bulk-price/workbook-preview",
        data={"workbook": (upload, "prices.xlsx"), "field": "sell_price"},
        content_type="multipart/form-data",
    ).get_json()

    assert preview["ok"]
    assert preview["count"] == 1
    assert preview["unchanged"] == 1
    assert preview["unmatched"] == 1
    assert preview["rows"][0]["label"] == "S-100"
    assert preview["rows"][0]["new_value"] == "125.00"

    applied = client.post("/parts-catalog/bulk-price/workbook-apply", json={
        "field": "sell_price",
        "source_name": "prices.xlsx",
        "updates": [{"id": ids["S-100"], "value": "125.00"}],
    }).get_json()
    assert applied["ok"] and applied["count"] == 1

    with get_session() as session:
        assert session.get(CatalogPart, ids["S-100"]).sell_price == Decimal("125.0000")
        assert session.get(CatalogPart, ids["S-200"]).sell_price == Decimal("57.5000")
        entry = session.query(PriceChangeLog).filter_by(change_type="upload").one()
        assert entry.old_value == Decimal("100.0000")
        assert entry.new_value == Decimal("125.0000")
        assert "prices.xlsx" in entry.batch_summary

    history = client.get("/parts-catalog/price-history").get_json()
    assert any(change["change_type"] == "upload" for change in history["changes"])


def test_workbook_upload_matches_on_part_name(tmp_path, monkeypatch):
    _, client, _ = _admin_client(tmp_path, monkeypatch)
    upload = _workbook(
        [("Buna Rod Seal 1.500", 12.25)],
        headers=("Part Name", "Sell Price"),
    )
    preview = client.post(
        "/parts-catalog/bulk-price/workbook-preview",
        data={"workbook": (upload, "by_name.xlsx"), "field": "sell_price"},
        content_type="multipart/form-data",
    ).get_json()
    assert preview["ok"] and preview["count"] == 1
    assert preview["rows"][0]["label"] == "S-300"


def test_workbook_upload_updates_unit_cost_and_sell_price(tmp_path, monkeypatch):
    _, client, ids = _admin_client(tmp_path, monkeypatch)
    upload = _workbook([
        ("S-100", 45.00, 125.00),
        ("S-200", 22.50, 60.00),
    ], headers=("Part Number", "Unit Cost", "Sell Price"))

    preview = client.post(
        "/parts-catalog/bulk-price/workbook-preview",
        data={"workbook": (upload, "prices.xlsx"), "field": "both"},
        content_type="multipart/form-data",
    ).get_json()

    assert preview["ok"] and preview["count"] == 2
    assert preview["rows"][0]["values"] == {"unit_cost": "45.00", "sell_price": "125.00"}

    applied = client.post("/parts-catalog/bulk-price/workbook-apply", json={
        "field": "both",
        "source_name": "prices.xlsx",
        "updates": [{"id": row["id"], "values": row["values"]} for row in preview["rows"]],
    }).get_json()
    assert applied["ok"] and applied["count"] == 4

    with get_session() as session:
        first = session.get(CatalogPart, ids["S-100"])
        second = session.get(CatalogPart, ids["S-200"])
        assert (first.unit_cost, first.sell_price) == (Decimal("45.0000"), Decimal("125.0000"))
        assert (second.unit_cost, second.sell_price) == (Decimal("22.5000"), Decimal("60.0000"))
        assert session.query(PriceChangeLog).filter_by(change_type="upload").count() == 4


def test_workbook_upload_rejects_unusable_file(tmp_path, monkeypatch):
    _, client, _ = _admin_client(tmp_path, monkeypatch)
    response = client.post(
        "/parts-catalog/bulk-price/workbook-preview",
        data={"workbook": (io.BytesIO(b"not a workbook"), "notes.txt"), "field": "sell_price"},
        content_type="multipart/form-data",
    )
    assert response.status_code == 400
    assert not response.get_json()["ok"]


def test_price_template_download_includes_vendor_column(tmp_path, monkeypatch):
    _, client, _ = _admin_client(tmp_path, monkeypatch)
    response = client.get("/parts-catalog/price-template.xlsx")
    assert response.status_code == 200

    from openpyxl import load_workbook

    sheet = load_workbook(io.BytesIO(response.data)).active
    headers = [cell.value for cell in sheet[1]]
    assert headers == ["Part Number", "Part Name", "Unit Cost", "Sell Price", "Vendors"]
    assert sheet.max_row == 4
