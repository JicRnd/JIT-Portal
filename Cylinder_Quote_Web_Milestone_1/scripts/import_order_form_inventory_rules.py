"""Import formula-preserving Order Form Inventory rows into Pricing.db.

This is an offline extraction tool. It uses pandas for the tabular Inventory
rows and openpyxl to preserve the original Excel formulas. Flask does not load
the workbook at runtime.
"""

from __future__ import annotations

import argparse
import re
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl
import pandas as pd
from sqlalchemy import delete

from app.db import get_session, init_db
from app.models_db import (
    OrderFormDependencyCell,
    OrderFormDependencyEdge,
    OrderFormInventoryRule,
    OrderFormWorkbookCell,
)


PROJECT_DIR = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = PROJECT_DIR / "2026 JIT Order Entry V3.xlsm"


def _audit_value(value):
    if isinstance(value, str):
        return value
    return None if value is None else str(value)


def _decimal(value):
    if value is None or pd.isna(value) or str(value).strip() == "":
        return None
    try:
        return Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, ValueError, TypeError):
        return None


def _formula_metadata(formula: str) -> tuple[str, str, str]:
    referenced_sheets = sorted(set(re.findall(r"([A-Za-z][A-Za-z0-9_]*)!", formula)))
    data_inputs = sorted(
        set(
            (column.upper(), row)
            for match in re.findall(
                r"(?:\[[^\]]+\])?Data!\$?([A-Z]{1,3})\$?(\d+)", formula, re.IGNORECASE
            )
            for column, row in (match,)
        )
    )
    if "H!" in formula.upper():
        category = "length_based"
    elif "ACC!" in formula.upper() or "F$26" in formula or "G$26" in formula:
        category = "helper_or_other"
    else:
        category = "quantity_or_option_based"
    return "|".join(referenced_sheets), "|".join(
        f"{column}{row}" for column, row in data_inputs
    ), category


def extract_rules(workbook_path: Path) -> list[dict]:
    workbook = openpyxl.load_workbook(
        workbook_path,
        keep_vba=workbook_path.suffix.lower() == ".xlsm",
        data_only=False,
        read_only=True,
    )
    try:
        worksheet = workbook["Inventory"]
        rows = []
        for row_number in range(2, 1030):
            rows.append(
                {
                    "source_row": row_number,
                    "part_number": worksheet.cell(row_number, 1).value,
                    "description": worksheet.cell(row_number, 2).value,
                    "unit_price": worksheet.cell(row_number, 3).value,
                    "allocation_formula": worksheet.cell(row_number, 5).value,
                }
            )
    finally:
        workbook.close()

    frame = pd.DataFrame(rows)
    frame["part_number"] = frame["part_number"].fillna("").astype(str).str.strip()
    frame["description"] = frame["description"].where(frame["description"].notna(), None)
    frame["allocation_formula"] = frame["allocation_formula"].where(
        frame["allocation_formula"].notna(), ""
    )
    frame = frame[
        (frame["part_number"] != "")
        & frame["allocation_formula"].astype(str).str.startswith("=")
    ]

    records = []
    for record in frame.to_dict(orient="records"):
        formula = str(record["allocation_formula"])
        referenced_sheets, data_inputs, category = _formula_metadata(formula)
        records.append(
            {
                "source_row": int(record["source_row"]),
                "part_number": record["part_number"],
                "description": record["description"],
                "unit_price": _decimal(record["unit_price"]),
                "allocation_formula": formula,
                "category": category,
                "referenced_sheets": referenced_sheets,
                "data_inputs": data_inputs,
                "source_workbook": workbook_path.name,
                "source_sheet": "Inventory",
            }
        )
    return records


def extract_helper_cells(workbook_path: Path) -> list[dict]:
    formulas = openpyxl.load_workbook(
        workbook_path,
        keep_vba=workbook_path.suffix.lower() == ".xlsm",
        data_only=False,
        read_only=True,
    )
    cached = openpyxl.load_workbook(
        workbook_path,
        keep_vba=workbook_path.suffix.lower() == ".xlsm",
        data_only=True,
        read_only=True,
    )
    rows = []
    try:
        sheet_names = {name.casefold(): name for name in formulas.sheetnames}
        for requested_name in ("H", "Acc", "Shipvia", "TieRod"):
            actual_name = sheet_names.get(requested_name.casefold())
            if actual_name is None:
                continue
            formula_sheet = formulas[actual_name]
            cached_sheet = cached[actual_name]
            for formula_row in formula_sheet.iter_rows():
                for formula_cell in formula_row:
                    if formula_cell.value is None:
                        continue
                    cached_value = cached_sheet[formula_cell.coordinate].value
                    formula = (
                        str(formula_cell.value)
                        if formula_cell.data_type == "f"
                        or (isinstance(formula_cell.value, str) and formula_cell.value.startswith("="))
                        else None
                    )
                    value = None if formula else formula_cell.value
                    rows.append(
                        {
                            "sheet_name": actual_name,
                            "coordinate": formula_cell.coordinate,
                            "formula": formula,
                            "value_text": str(value) if isinstance(value, str) else None,
                            "value_number": _decimal(value) if not isinstance(value, str) else None,
                            "source_workbook": workbook_path.name,
                        }
                    )
    finally:
        formulas.close()
        cached.close()
    return rows


def extract_dependency_audit(audit_path: Path) -> tuple[list[dict], list[dict]]:
    workbook = openpyxl.load_workbook(audit_path, data_only=False, read_only=True)
    try:
        cell_sheet = workbook["Dependency cells"]
        edge_sheet = workbook["Dependency edges"]
        cells = []
        for row in cell_sheet.iter_rows(min_row=6, values_only=True):
            if not row[0]:
                continue
            cells.append(
                {
                    "cell_key": str(row[0]),
                    "kind": str(row[1] or ""),
                    "formula": row[2],
                    "value_text": row[3] if isinstance(row[3], str) else None,
                    "value_number": _decimal(row[3]) if not isinstance(row[3], str) else None,
                    "cached_value": _audit_value(row[4]),
                    "direct_precedents": row[5],
                    "used_by_part_formulas": row[6],
                    "issues": row[7],
                    "source_workbook": audit_path.name,
                }
            )
        edges = []
        for row in edge_sheet.iter_rows(min_row=6, values_only=True):
            if not row[0] or not row[1]:
                continue
            edges.append(
                {
                    "from_cell": str(row[0]),
                    "to_reference": str(row[1]),
                    "reference_kind": row[2],
                    "formula_token": row[3],
                    "source_workbook": audit_path.name,
                }
            )
        return cells, edges
    finally:
        workbook.close()


def import_rules(workbook_path: Path, audit_path: Path | None = None) -> int:
    rules = extract_rules(workbook_path)
    helper_cells = extract_helper_cells(workbook_path)
    dependency_cells, dependency_edges = (
        extract_dependency_audit(audit_path) if audit_path else ([], [])
    )
    init_db()
    with get_session() as session:
        session.execute(delete(OrderFormInventoryRule))
        session.execute(delete(OrderFormWorkbookCell))
        session.execute(delete(OrderFormDependencyCell))
        session.execute(delete(OrderFormDependencyEdge))
        session.add_all([OrderFormInventoryRule(**rule) for rule in rules])
        session.add_all([OrderFormWorkbookCell(**cell) for cell in helper_cells])
        session.add_all([OrderFormDependencyCell(**cell) for cell in dependency_cells])
        session.add_all([OrderFormDependencyEdge(**edge) for edge in dependency_edges])
        session.commit()
    return len(rules)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", nargs="?", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--audit", type=Path, default=None)
    args = parser.parse_args()
    count = import_rules(args.workbook, args.audit)
    print(f"Imported {count} Order Form Inventory rules from {args.workbook.name}.")


if __name__ == "__main__":
    main()
