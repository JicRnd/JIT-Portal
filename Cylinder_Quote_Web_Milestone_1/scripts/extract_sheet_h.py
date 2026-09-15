"""Extract Sheet H cells into a formula-preserving JSON table.

This is an offline analysis tool. It deliberately does not evaluate Excel
formulas or become part of the Flask runtime.
"""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd


PROJECT_DIR = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = PROJECT_DIR / "2026 JIT Order Entry V3.xlsm"
DEFAULT_OUTPUT = PROJECT_DIR / "data" / "sheet_h_cells.json"
DEFAULT_ENGINEERING_OUTPUT = PROJECT_DIR / "app" / "static" / "sheet_h_engineering.json"
WEIGHT_COLUMNS = ("bore", "rod", "weight_base", "weight_per_stroke", "unused", "formula")

SEMANTIC_COLUMNS = {
    "M": "bore",
    "N": "rod",
    "O": "rod_thread_style_1",
    "P": "rod_thread_style_2",
    "Z": "non_cushion_base",
    "AA": "cushion_base",
    "AB": "weight_per_stroke_or_rate",
    "AL": "rod_or_mount_constant",
}

ROD_MATRIX_COLUMNS = {
    "AM": ("A", "LH", "SA", ("BE", "CE"), (4,)),
    "AN": ("A", "LH", "SA", ("RE", "NC"), (4,)),
    "AO": ("A", "LH", "SA", ("BE", "CE"), (3,)),
    "AP": ("A", "LH", "SA", ("RE", "NC"), (3,)),
    "AQ": ("A", "LH", "SA", ("BE", "CE"), (1, 2, 6)),
    "AR": ("A", "LH", "SA", ("RE", "NC"), (1, 2, 6)),
    "AV": ("H", "HM", ("BE", "CE"), (4,)),
    "AW": ("H", "HM", ("RE", "NC"), (4,)),
    "AX": ("H", "HM", ("BE", "CE"), (3,)),
    "AY": ("H", "HM", ("RE", "NC"), (3,)),
    "AZ": ("H", "HM", ("BE", "CE"), (1, 2, 6)),
    "BA": ("H", "HM", ("RE", "NC"), (1, 2, 6)),
}


def extract_sheet_cells(workbook_path: Path, sheet_name: str = "H") -> dict[str, Any]:
    """Return non-empty cells from a workbook sheet as a JSON-safe table."""
    workbook = openpyxl.load_workbook(
        workbook_path,
        keep_vba=workbook_path.suffix.lower() == ".xlsm",
        data_only=False,
        read_only=True,
    )
    try:
        worksheet = workbook[sheet_name]
        rows = []
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                is_formula = cell.data_type == "f" or (
                    isinstance(cell.value, str) and cell.value.startswith("=")
                )
                rows.append(
                    {
                        "coordinate": cell.coordinate,
                        "row": cell.row,
                        "column": cell.column,
                        "value": cell.value,
                        "is_formula": is_formula,
                    }
                )
    finally:
        workbook.close()

    table = pd.DataFrame(
        rows,
        columns=["coordinate", "row", "column", "value", "is_formula"],
    )
    records = table.where(table.notna(), None).to_dict(orient="records")
    selector_rows = build_selector_rows(table)
    rod_addition_rules = build_rod_addition_rules(table)
    barrel_rules = build_barrel_rules(table)
    torque_rules = extract_torque_rules(workbook_path)
    return {
        "source_workbook": workbook_path.name,
        "source_sheet": sheet_name,
        "table_type": "non_empty_cells",
        "row_count": int(table["row"].nunique()) if not table.empty else 0,
        "cell_count": int(len(table)),
        "cells": records,
        "selector_rows": selector_rows,
        "rod_addition_rules": rod_addition_rules,
        "barrel_rules": barrel_rules,
        "torque_rules": torque_rules,
    }


def build_selector_rows(table: pd.DataFrame) -> list[dict[str, Any]]:
    """Build a semantic view of H rows without evaluating Excel formulas."""
    if table.empty:
        return []

    by_row = table.set_index(["row", "column"])["value"]
    rows = []
    for row_number in sorted(table["row"].unique()):
        values = {
            column: by_row.get((row_number, index), None)
            for index, column in enumerate(
                ["A", "B", "C", "D", "E", "F", "G", "H", "I", "J", "K", "L", "M", "N", "O", "P", "Q", "R", "S", "T", "U", "V", "W", "X", "Y", "Z", "AA", "AB", "AC", "AD", "AE", "AF", "AG", "AH", "AI", "AJ", "AK", "AL", "AM", "AN", "AO", "AP", "AQ", "AR", "AS", "AT", "AU", "AV", "AW", "AX", "AY", "AZ", "BA", "BB"],
                start=1,
            )
        }
        if values["M"] is None or values["N"] is None:
            continue
        row_data = {
            "source_row": int(row_number),
            "bore": values["M"],
            "rod": values["N"],
            "source_values": {
                column: values[column]
                for column in values
                if values[column] is not None
            },
        }
        for column, name in SEMANTIC_COLUMNS.items():
            row_data[name] = values[column]
        rows.append(row_data)
    return rows


def build_rod_addition_rules(table: pd.DataFrame) -> list[dict[str, Any]]:
    """Resolve H rod-matrix selectors to numeric helper values."""
    if table.empty:
        return []

    values = {
        str(record["coordinate"]): record["value"]
        for record in table.to_dict(orient="records")
    }
    rules = []
    for record in table.to_dict(orient="records"):
        column = next(
            (
                letter
                for letter in ROD_MATRIX_COLUMNS
                if record["coordinate"].startswith(letter)
            ),
            None,
        )
        formula = record["value"]
        if column is None or not isinstance(formula, str) or not formula.startswith("="):
            continue
        bore_match = re.search(r"Data!\$?B\$?4\s*=\s*([0-9.]+)", formula)
        rod_match = re.search(r"Data!\$?B\$?6\s*=\s*([0-9.]+)", formula)
        helper_match = re.search(r",\s*([A-Z]{2}[0-9]+),\s*\"\"", formula)
        if not bore_match or not rod_match or not helper_match:
            continue
        helper_coordinate = helper_match.group(1)
        helper_value = values.get(helper_coordinate)
        if not isinstance(helper_value, (int, float)):
            continue
        selector = ROD_MATRIX_COLUMNS[column]
        if column in {"AM", "AN", "AO", "AP", "AQ", "AR"}:
            families = selector[0:3]
            cushions = selector[3]
            styles = selector[4]
        else:
            families = selector[0:2]
            cushions = selector[2]
            styles = selector[3]
        rules.append(
            {
                "source_row": int(record["row"]),
                "source_column": column,
                "helper_cell": helper_coordinate,
                "families": list(families),
                "bore": float(bore_match.group(1)),
                "rod": float(rod_match.group(1)),
                "cushions": list(cushions),
                "rod_styles": list(styles),
                "rod_length_addition": helper_value,
            }
        )
    return rules


def build_barrel_rules(table: pd.DataFrame) -> list[dict[str, Any]]:
    """Extract the Z:AJ bore-specific cap/barrel additions."""
    if table.empty:
        return []
    by_coordinate = {
        str(record["coordinate"]): record["value"]
        for record in table.to_dict(orient="records")
    }
    rules = []
    for row_number in range(7, 19):
        bore = by_coordinate.get(f"AG{row_number}")
        non_h_addition = by_coordinate.get(f"AH{row_number}")
        h_addition = by_coordinate.get(f"AI{row_number}")
        selector_formula = by_coordinate.get(f"AJ{row_number}")
        if not all(isinstance(value, (int, float)) for value in (bore, non_h_addition, h_addition)):
            continue
        rules.append(
            {
                "source_row": row_number,
                "bore": bore,
                "families": {
                    "A": non_h_addition,
                    "LH": non_h_addition,
                    "SA": non_h_addition,
                    "H": h_addition,
                    "HM": h_addition,
                },
                "non_h_addition": non_h_addition,
                "h_addition": h_addition,
                "selector_formula": selector_formula,
            }
        )
    return rules


def extract_torque_rules(workbook_path: Path) -> list[dict[str, Any]]:
    """Extract the Order Form torque lookup used by AN1."""
    workbook = openpyxl.load_workbook(
        workbook_path,
        keep_vba=workbook_path.suffix.lower() == ".xlsm",
        data_only=False,
        read_only=True,
    )
    try:
        worksheet = workbook["Order Form"]
        rules = []
        for row_number in range(2, 12):
            bore = worksheet[f"AK{row_number}"].value
            a_torque = worksheet[f"AL{row_number}"].value
            h_torque = worksheet[f"AM{row_number}"].value
            if isinstance(bore, (int, float)):
                rules.append(
                    {
                        "source_row": row_number,
                        "bore": bore,
                        "families": {"A": a_torque, "LH": a_torque, "H": h_torque, "HM": h_torque},
                    }
                )
        return rules
    finally:
        workbook.close()


def extract_weight_rules(workbook_path: Path) -> list[dict[str, Any]]:
    """Extract the A/LH cylinder weight chart from Data!Y:AD with pandas."""
    table = pd.read_excel(
        workbook_path,
        sheet_name="Data",
        usecols="Y:AD",
        header=None,
        engine="openpyxl",
    )
    table.columns = WEIGHT_COLUMNS
    table["bore"] = table["bore"].ffill()
    rules = []
    for row_number, row in table.iterrows():
        if not all(pd.notna(row[column]) for column in WEIGHT_COLUMNS[:4]):
            continue
        if not all(pd.api.types.is_number(row[column]) for column in WEIGHT_COLUMNS[:4]):
            continue
        rules.append(
            {
                "source_sheet": "Data",
                "source_row": int(row_number) + 1,
                "source_range": "Y:AD",
                "families": ["A", "LH"],
                "bore": float(row["bore"]),
                "rod": float(row["rod"]),
                "weight_base": float(row["weight_base"]),
                "weight_per_stroke": float(row["weight_per_stroke"]),
                "calculation": "weight_base + weight_per_stroke * stroke",
                "selector_formula": row["formula"] if pd.notna(row["formula"]) else None,
            }
        )
    return rules


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Extract non-empty Sheet H cells into a JSON table."
    )
    parser.add_argument(
        "--workbook",
        type=Path,
        default=DEFAULT_WORKBOOK,
        help="Path to the source XLSX/XLSM workbook.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=DEFAULT_OUTPUT,
        help="Path for the extracted JSON table.",
    )
    parser.add_argument(
        "--engineering-output",
        type=Path,
        default=DEFAULT_ENGINEERING_OUTPUT,
        help="Path for the compact runtime engineering table.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    if not args.workbook.exists():
        raise SystemExit(f"Workbook not found: {args.workbook}")

    extracted = extract_sheet_cells(args.workbook)
    weight_rules = extract_weight_rules(args.workbook)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(extracted, indent=2), encoding="utf-8")
    engineering = {
        "source_workbook": extracted["source_workbook"],
        "source_sheet": extracted["source_sheet"],
        "purpose": "Order Form testing-section engineering lookups",
        "rod_addition_rules": extracted["rod_addition_rules"],
        "barrel_rules": extracted["barrel_rules"],
        "torque_rules": extracted["torque_rules"],
        "weight_rules": weight_rules,
    }
    args.engineering_output.parent.mkdir(parents=True, exist_ok=True)
    args.engineering_output.write_text(json.dumps(engineering, indent=2), encoding="utf-8")
    print(
        f"Extracted {extracted['cell_count']} non-empty cells from "
        f"{extracted['source_sheet']} to {args.output}; "
        f"wrote {len(extracted['rod_addition_rules'])} engineering rules to "
        f"{args.engineering_output}"
    )


if __name__ == "__main__":
    main()
