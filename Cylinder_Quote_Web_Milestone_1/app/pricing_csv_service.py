from __future__ import annotations

import csv
import os
import tempfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl

from cylinder_quote_engine.data import PricingData


PRICING_DATA_DIRECTORY = Path(__file__).resolve().parents[1] / "data"
PRICING_MIRROR_DIRECTORY = Path(__file__).resolve().parent / "Pricing" / "pricing_file_mirrors"
PRICING_HISTORY_PATH = Path(__file__).resolve().parent / "Pricing" / "pricing_calcuator_change_history.txt"
_EXCLUDED_PRICE_HEADERS = {
    "rate_name",
    "rate_1_name",
    "rate_2_name",
    "rate_1_col",
    "rate_2_col",
}
_EDITABLE_SUFFIXES = ("_per_in", "_per_stroke", "_per_length", "_per_end", "_per_unit")
_RUNTIME_SERIES_PRICING_FILES = (
    {
        "filename": "standard_series_pricing.xlsx",
        "row_count": 458,
        "editable_columns": ["stroke_rate", "cushion_per_end", "base_price"],
        "description": "Editable mirror; protected runtime workbook remains unchanged.",
    },
    {
        "filename": "metric_series_pricing.xlsx",
        "row_count": 156,
        "editable_columns": ["stroke_rate", "cushion_per_end", "base_price"],
        "description": "Editable mirror; protected runtime workbook remains unchanged.",
    },
)
_RUNTIME_SERIES_PRICING_NAMES = {file["filename"] for file in _RUNTIME_SERIES_PRICING_FILES}
_RETIRED_SERIES_PRICING_XLSX = "series_pricing.xlsx"


class PricingCsvError(ValueError):
    def __init__(self, message: str, status_code: int = 400):
        super().__init__(message)
        self.status_code = status_code


def _csv_path(filename: str) -> Path:
    candidate = Path(filename or "")
    if candidate.name != filename or candidate.suffix.lower() != ".csv":
        raise PricingCsvError("Choose a valid pricing CSV file.")
    path = PRICING_DATA_DIRECTORY / candidate.name
    if not path.is_file():
        legacy_path = PRICING_DATA_DIRECTORY.parent / "Unused_files" / candidate.name
        if legacy_path.is_file():
            return legacy_path
        raise PricingCsvError("That pricing CSV file was not found.", 404)
    return path


def _active_xlsx_path(filename: str) -> Path:
    candidate = Path(filename or "")
    if candidate.name != filename or candidate.suffix.lower() != ".xlsx":
        raise PricingCsvError("Choose a valid pricing workbook.")
    if candidate.name == _RETIRED_SERIES_PRICING_XLSX:
        raise PricingCsvError(
            "The combined series pricing workbook is retired; choose "
            "standard_series_pricing.xlsx or metric_series_pricing.xlsx."
        )
    path = PRICING_DATA_DIRECTORY / candidate.name
    if not path.is_file():
        raise PricingCsvError("That pricing workbook was not found.", 404)
    return path


def _xlsx_mirror_path(filename: str) -> Path:
    candidate = Path(filename or "")
    if candidate.name != filename or candidate.name not in _RUNTIME_SERIES_PRICING_NAMES:
        raise PricingCsvError("Choose a valid pricing workbook.")
    path = PRICING_MIRROR_DIRECTORY / candidate.name
    if not path.is_file():
        raise PricingCsvError("That pricing workbook mirror was not found.", 404)
    return path


def _is_editable_header(header: str) -> bool:
    normalized = header.strip().lower()
    if normalized in _EXCLUDED_PRICE_HEADERS or normalized.endswith("_source") or normalized.endswith("_formula"):
        return False
    return (
        normalized == "price"
        or "price" in normalized
        or "rate" in normalized
        or "cost" in normalized
        or normalized.endswith(_EDITABLE_SUFFIXES)
    )


def _read_csv(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    with path.open("r", encoding="utf-8-sig", newline="") as source:
        reader = csv.DictReader(source)
        headers = reader.fieldnames or []
        if not headers or any(not header or not header.strip() for header in headers):
            raise PricingCsvError(f"{path.name} has an invalid header row.")
        rows = list(reader)
    return headers, rows


def _display_columns(filename: str, headers: list[str]) -> tuple[list[str], dict[str, str]]:
    if filename == "rod_boot_rates.csv":
        visible = [
            header
            for header in headers
            if "source" not in header.lower() and "formula" not in header.lower()
        ]
        labels = {header: header for header in visible}
        labels["per_stroke_rate"] = '$/"'
        return visible, labels

    if filename not in {"series_pricing.csv", *_RUNTIME_SERIES_PRICING_NAMES}:
        return headers, {header: header for header in headers}

    hidden = {
        header for header in headers
        if header == "stroke_unit" or "source" in header.lower() or "formula" in header.lower()
    }
    visible = [header for header in headers if header not in hidden]
    visible.remove("mount_codes")
    visible.remove("base_price")
    aliases_index = visible.index("series_aliases") + 1
    visible.insert(aliases_index, "mount_codes")
    cushion_index = visible.index("cushion_per_end") + 1
    visible.insert(cushion_index, "base_price")
    labels = {header: header for header in visible}
    labels["stroke_rate"] = "Stroke rate per in"
    return visible, labels


def list_pricing_csvs() -> list[dict[str, object]]:
    files = []
    for path in sorted(PRICING_DATA_DIRECTORY.glob("*.xlsx")):
        if path.name.startswith("~$"):
            continue
        if path.name in {
            _RETIRED_SERIES_PRICING_XLSX,
            "series_row_rates.xlsx",
            *_RUNTIME_SERIES_PRICING_NAMES,
        }:
            continue
        headers, rows = _read_xlsx(path)
        editable_columns = [header for header in headers if _is_editable_header(header)]
        if not editable_columns:
            continue
        files.append({
            "filename": path.name,
            "row_count": len(rows),
            "editable_columns": editable_columns,
        })
    return files


def list_runtime_series_pricing_files() -> list[dict[str, object]]:
    return [dict(file) for file in _RUNTIME_SERIES_PRICING_FILES]


def _pricing_document(
    filename: str,
    headers: list[str],
    rows: list[dict[str, str]],
    is_mirror: bool = False,
) -> dict[str, object]:
    editable_columns = [header for header in headers if _is_editable_header(header)]
    if not editable_columns:
        raise PricingCsvError("That file has no editable pricing columns.", 400)
    display_headers, display_labels = _display_columns(filename, headers)
    return {
        "filename": filename,
        "headers": headers,
        "display_headers": display_headers,
        "display_labels": display_labels,
        "rows": rows,
        "editable_columns": editable_columns,
        "is_mirror": is_mirror,
    }


def read_pricing_csv(filename: str) -> dict[str, object]:
    path = _csv_path(filename)
    headers, rows = _read_csv(path)
    return _pricing_document(path.name, headers, rows)


def _xlsx_cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, date):
        return f"{value.month}/{value.day}-{value.year % 100:02d}"
    return str(value)


def _read_xlsx(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    workbook = openpyxl.load_workbook(path, data_only=False, read_only=True)
    try:
        worksheet = workbook["series_pricing"] if "series_pricing" in workbook.sheetnames else workbook.active
        header_values = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
        headers = [_xlsx_cell_text(value) for value in (header_values or ())]
        if not headers or any(not header.strip() for header in headers):
            raise PricingCsvError(f"{path.name} has an invalid header row.")
        rows = []
        for values in worksheet.iter_rows(min_row=2, values_only=True):
            rows.append({header: _xlsx_cell_text(value) for header, value in zip(headers, values)})
        return headers, rows
    finally:
        workbook.close()


def read_pricing_xlsx(filename: str) -> dict[str, object]:
    is_mirror = filename in _RUNTIME_SERIES_PRICING_NAMES
    path = _xlsx_mirror_path(filename) if is_mirror else _active_xlsx_path(filename)
    headers, rows = _read_xlsx(path)
    return _pricing_document(path.name, headers, rows, is_mirror=is_mirror)


def read_pricing_file(filename: str) -> dict[str, object]:
    if str(filename or "").lower().endswith(".xlsx"):
        return read_pricing_xlsx(filename)
    return read_pricing_csv(filename)


def _validate_numeric(value: str, filename: str, header: str, row_number: int) -> str:
    value = str(value or "").strip()
    if not value:
        return ""
    try:
        amount = Decimal(value)
    except InvalidOperation as exc:
        raise PricingCsvError(f"{filename} row {row_number}: {header} must be numeric.") from exc
    if not amount.is_finite() or amount < 0:
        raise PricingCsvError(f"{filename} row {row_number}: {header} must be zero or greater.")
    return value


def save_pricing_csv(filename: str, headers: list[str], rows: list[dict[str, object]], actor: str) -> dict[str, object]:
    path = _csv_path(filename)
    current_headers, current_rows = _read_csv(path)
    if headers != current_headers:
        raise PricingCsvError("The pricing CSV changed while it was open. Reload it before saving.", 409)
    if len(rows) != len(current_rows):
        raise PricingCsvError("Rows cannot be added or removed from this editor.", 400)

    editable_columns = [header for header in headers if _is_editable_header(header)]
    normalized_rows: list[dict[str, str]] = []
    changed_cells = 0
    for row_number, row in enumerate(rows, start=2):
        if set(row) != set(headers):
            raise PricingCsvError(f"Row {row_number} does not match the CSV headers.")
        normalized: dict[str, str] = {}
        for header in headers:
            value = str(row.get(header) or "").strip()
            if header in editable_columns:
                value = _validate_numeric(value, path.name, header, row_number)
                if value != current_rows[row_number - 2].get(header, "").strip():
                    changed_cells += 1
            elif value != current_rows[row_number - 2].get(header, "").strip():
                raise PricingCsvError(f"Only pricing columns can be edited: {header}.")
            normalized[header] = value
        normalized_rows.append(normalized)

    if changed_cells:
        fd, temporary_name = tempfile.mkstemp(prefix=f".{path.stem}.", suffix=".csv", dir=path.parent)
        try:
            with os.fdopen(fd, "w", encoding="utf-8", newline="") as target:
                writer = csv.DictWriter(target, fieldnames=headers, lineterminator="\n")
                writer.writeheader()
                writer.writerows(normalized_rows)
                target.flush()
                os.fsync(target.fileno())
            os.replace(temporary_name, path)
        finally:
            if os.path.exists(temporary_name):
                os.unlink(temporary_name)
        _reload_runtime_pricing()
        _append_history(actor, path.name, changed_cells)

    return {"filename": path.name, "changed_cells": changed_cells}


def _xlsx_numeric_value(value: str) -> int | float | None:
    if not value:
        return None
    amount = Decimal(value)
    return int(amount) if amount == amount.to_integral_value() else float(amount)


def _write_xlsx(path: Path, headers: list[str], rows: list[dict[str, str]], editable_columns: list[str]) -> None:
    workbook = openpyxl.load_workbook(path)
    temporary_name = ""
    try:
        worksheet = workbook["series_pricing"] if "series_pricing" in workbook.sheetnames else workbook.active
        column_numbers = {header: index + 1 for index, header in enumerate(headers)}
        for row_number, row in enumerate(rows, start=2):
            for header in editable_columns:
                worksheet.cell(row=row_number, column=column_numbers[header]).value = _xlsx_numeric_value(row[header])
        fd, temporary_name = tempfile.mkstemp(prefix=f".{path.stem}.", suffix=".xlsx", dir=path.parent)
        os.close(fd)
        workbook.save(temporary_name)
        os.replace(temporary_name, path)
    finally:
        workbook.close()
        if temporary_name and os.path.exists(temporary_name):
            os.unlink(temporary_name)


def save_pricing_xlsx(filename: str, headers: list[str], rows: list[dict[str, object]], actor: str) -> dict[str, object]:
    is_mirror = filename in _RUNTIME_SERIES_PRICING_NAMES
    path = _xlsx_mirror_path(filename) if is_mirror else _active_xlsx_path(filename)
    current_headers, current_rows = _read_xlsx(path)
    if headers != current_headers:
        raise PricingCsvError("The pricing workbook changed while it was open. Reload it before saving.", 409)
    if len(rows) != len(current_rows):
        raise PricingCsvError("Rows cannot be added or removed from this editor.", 400)

    editable_columns = [header for header in headers if _is_editable_header(header)]
    normalized_rows: list[dict[str, str]] = []
    changed_cells = 0
    for row_number, row in enumerate(rows, start=2):
        if set(row) != set(headers):
            raise PricingCsvError(f"Row {row_number} does not match the workbook headers.")
        normalized: dict[str, str] = {}
        for header in headers:
            value = str(row.get(header) or "").strip()
            if header in editable_columns:
                value = _validate_numeric(value, path.name, header, row_number)
                if value != current_rows[row_number - 2].get(header, "").strip():
                    changed_cells += 1
            elif value != current_rows[row_number - 2].get(header, "").strip():
                raise PricingCsvError(f"Only pricing columns can be edited: {header}.")
            normalized[header] = value
        normalized_rows.append(normalized)

    if changed_cells:
        _write_xlsx(path, headers, normalized_rows, editable_columns)
        _append_history(actor, path.name, changed_cells, "XLSX mirror" if is_mirror else "XLSX")

    return {"filename": path.name, "changed_cells": changed_cells}


def _reload_runtime_pricing() -> None:
    from . import quote_service
    from flask import current_app, has_app_context

    runtime_engine = current_app.extensions.get("pricing_engine") if has_app_context() else None
    if runtime_engine is not None:
        runtime_engine.data = PricingData(PRICING_DATA_DIRECTORY)
    quote_service._ACCESSORY_ENGINE.data = PricingData(PRICING_DATA_DIRECTORY)


def _append_history(actor: str, filename: str, changed_cells: int, source_type: str = "CSV") -> None:
    timestamp = datetime.now().isoformat(timespec="seconds")
    with PRICING_HISTORY_PATH.open("a", encoding="utf-8") as history:
        history.write(
            f"\n{timestamp} - Admin {source_type} pricing update - file={filename}; "
            f"changed_cells={changed_cells}; actor={actor}\n"
        )
