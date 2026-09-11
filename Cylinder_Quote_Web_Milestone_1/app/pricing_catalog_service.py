from __future__ import annotations

import csv
import re
import uuid
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path

from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from .db import get_session, init_db
from .models_db import (
    BaseAssemblyPrice,
    CatalogPart,
    CommonModificationPrice,
    PhVaPrice,
    PriceChangeLog,
    utc_now,
)
from .part_family_service import (
    SIZE_RULES,
    TEMPLATE_VARIABLES,
    describe_part,
    family_headline,
    list_families,
)


PRICING_DATA_DIRECTORY = Path(__file__).resolve().parents[1] / "pricing_data"
MONEY_QUANT = Decimal("0.01")
STORAGE_QUANT = Decimal("0.0001")


class PriceUpdateError(ValueError):
    def __init__(self, message: str, status_code: int = 400):
        super().__init__(message)
        self.status_code = status_code


EDITABLE_PRICE_FIELDS = {
    "catalog_parts": (CatalogPart, {"unit_cost", "sell_price"}),
    "base_assembly_prices": (
        BaseAssemblyPrice,
        {"base_price", "cushion_base", "per_unit_price", "cushion_per_end"},
    ),
    "common_modification_prices": (CommonModificationPrice, {"base_price", "per_unit_price"}),
    "ph_va_prices": (PhVaPrice, {"base_price", "per_unit_price"}),
}


def _clean(value: str | None) -> str | None:
    value = (value or "").strip()
    return value or None


def _decimal(value: str | None) -> Decimal | None:
    value = _clean(value)
    if value is None:
        return None
    try:
        return Decimal(value)
    except InvalidOperation:
        return None


def _source_row(value: str | None) -> int:
    try:
        return int((value or "").strip())
    except ValueError as exc:
        raise ValueError(f"Invalid source row: {value}") from exc


def _raw_price(value: str | None) -> str | None:
    value = _clean(value)
    return value if value and _decimal(value) is None else None


def _unit_label(unit: str | None) -> str:
    value = (unit or "").strip().lower()
    if value in {"in", "inch", "inches"}:
        return "in"
    if value in {"mm", "millimeter", "millimeters"}:
        return "mm"
    return value


def format_money(value) -> str:
    if value in (None, ""):
        return ""
    try:
        amount = Decimal(str(value)).quantize(MONEY_QUANT)
    except (InvalidOperation, ValueError):
        return str(value)
    return f"${amount:,.2f}"


def format_price_unit(unit: str | None) -> str:
    unit_label = _unit_label(unit)
    return f" / {unit_label}" if unit_label in {"in", "mm"} else ""


def format_measurement(value, unit: str | None) -> str:
    if value in (None, ""):
        return ""
    unit_label = _unit_label(unit)
    if unit_label not in {"in", "mm"}:
        return str(value)
    text = str(value).strip()
    try:
        amount = Decimal(text)
    except InvalidOperation:
        return text
    if unit_label == "mm":
        number = f"{amount:f}".rstrip("0").rstrip(".")
    else:
        decimal_places = max(2, min(4, -amount.as_tuple().exponent))
        number = f"{amount:.{decimal_places}f}".rstrip("0").rstrip(".")
        if "." not in number:
            number = f"{number}.00"
        elif len(number.rsplit(".", 1)[1]) == 1:
            number = f"{number}0"
    return f"{number} {unit_label}"


def _csv_rows(filename: str) -> list[dict[str, str]]:
    path = PRICING_DATA_DIRECTORY / filename
    with path.open("r", encoding="utf-8-sig", newline="") as source:
        return list(csv.DictReader(source))


def _upsert_by(session: Session, model, filters: dict, values: dict) -> None:
    row = session.execute(select(model).filter_by(**filters)).scalar_one_or_none()
    if row is None:
        session.add(model(**values))
        return
    for field, value in values.items():
        setattr(row, field, value)


def import_pricing_catalog() -> dict[str, int]:
    """Import the four extracted CSVs into Pricing.db without duplicate rows."""
    init_db()
    counts = {
        "catalog_parts": 0,
        "base_assembly_prices": 0,
        "common_modification_prices": 0,
        "ph_va_prices": 0,
    }

    with get_session() as session:
        for row in _csv_rows("jit_master_parts_catalog.csv"):
            part_number = _clean(row.get("part_number"))
            if not part_number:
                continue
            values = {
                "part_number": part_number,
                "description": _clean(row.get("description")),
                "unit_cost": _decimal(row.get("unit_cost")),
                "sell_price": _decimal(row.get("sell_price")),
                "sell_price_source": _clean(row.get("sell_price_source")),
                "inventory_cost_values": _clean(row.get("inventory_cost_values")),
                "shipvia_sell_values": _clean(row.get("shipvia_sell_values")),
                "acc_sell_values": _clean(row.get("acc_sell_values")),
                "all_descriptions": _clean(row.get("all_descriptions")),
                "source_locations": _clean(row.get("source_locations")),
                "review_needed": _clean(row.get("review_needed")),
                "review_reason": _clean(row.get("review_reason")),
                "active": _clean(row.get("active")),
            }
            _upsert_by(session, CatalogPart, {"part_number": part_number}, values)
            counts["catalog_parts"] += 1

        for source_import_index, row in enumerate(_csv_rows("jit_base_and_assembly_pricing_clean.csv"), start=1):
            values = {
                "pricing_type": _clean(row.get("pricing_type")) or "unknown",
                "series": _clean(row.get("series")),
                "bore": _clean(row.get("bore")),
                "rod": _clean(row.get("rod")),
                "mount_group": _clean(row.get("mount_group")),
                "base_price": _decimal(row.get("base_price")),
                "base_price_raw": _raw_price(row.get("base_price")),
                "cushion_base": _decimal(row.get("cushion_base")),
                "cushion_base_raw": _raw_price(row.get("cushion_base")),
                "per_unit_price": _decimal(row.get("per_unit_price")),
                "per_unit_price_raw": _raw_price(row.get("per_unit_price")),
                "cushion_per_end": _decimal(row.get("cushion_per_end")),
                "cushion_per_end_raw": _raw_price(row.get("cushion_per_end")),
                "rod_thread_std": _clean(row.get("rod_thread_std")),
                "rod_thread_oversize": _clean(row.get("rod_thread_oversize")),
                "unit": _clean(row.get("unit")),
                "source_sheet": _clean(row.get("source_sheet")) or "unknown",
                "source_row": _source_row(row.get("source_row")),
                "source_import_index": source_import_index,
            }
            _upsert_by(session, BaseAssemblyPrice, {"source_import_index": source_import_index}, values)
            counts["base_assembly_prices"] += 1

        for source_import_index, row in enumerate(_csv_rows("jit_common_modification_pricing_corrected.csv"), start=1):
            values = {
                "series": _clean(row.get("series")),
                "bore": _clean(row.get("bore")),
                "rod": _clean(row.get("rod")),
                "option_name": _clean(row.get("option_name")) or "unknown",
                "base_price": _decimal(row.get("base_price")),
                "base_price_raw": _raw_price(row.get("base_price")),
                "per_unit_price": _decimal(row.get("per_unit_price")),
                "per_unit_price_raw": _raw_price(row.get("per_unit_price")),
                "unit": _clean(row.get("unit")),
                "source_sheet": _clean(row.get("source_sheet")) or "unknown",
                "source_row": _source_row(row.get("source_row")),
                "source_import_index": source_import_index,
            }
            _upsert_by(session, CommonModificationPrice, {"source_import_index": source_import_index}, values)
            counts["common_modification_prices"] += 1

        for source_import_index, row in enumerate(_csv_rows("jit_ph_va_pricing.csv"), start=1):
            values = {
                "pricing_group": _clean(row.get("pricing_group")) or "unknown",
                "series": _clean(row.get("series")),
                "bore": _clean(row.get("bore")),
                "rod": _clean(row.get("rod")),
                "option_name": _clean(row.get("option_name")) or "unknown",
                "base_price": _decimal(row.get("base_price")),
                "base_price_raw": _raw_price(row.get("base_price")),
                "per_unit_price": _decimal(row.get("per_unit_price")),
                "per_unit_price_raw": _raw_price(row.get("per_unit_price")),
                "unit": _clean(row.get("unit")),
                "source_sheet": _clean(row.get("source_sheet")) or "unknown",
                "source_row": _source_row(row.get("source_row")),
                "source_import_index": source_import_index,
            }
            _upsert_by(session, PhVaPrice, {"source_import_index": source_import_index}, values)
            counts["ph_va_prices"] += 1

        session.commit()
    return counts


# --- Traditional catalog categories -------------------------------------


# Preferred display/index order. Anything left unmatched falls to the last entry.
CATEGORY_ORDER = [
    "Seal Kits",
    "Seals",
    "Wipers",
    "Pistons",
    "Glands",
    "Barrels",
    "Rods",
    "Tie Rods",
    "Jam Nuts",
    "Mounts / Mounting Hardware",
    "Rod Eyes / Clevises",
    "Bearings / Bushings",
    "Fasteners",
    "Accessories",
    "Other / Unclassified",
]

# Checked in order (most specific first) against the part's description text.
_CATEGORY_KEYWORDS = [
    ("Seal Kits", ("seal kit", "seal kits")),
    ("Wipers", ("wiper", "scraper")),
    ("Glands", ("gland",)),
    ("Tie Rods", ("tie rod", "tie-rod", "tierod")),
    ("Rod Eyes / Clevises", ("rod eye", "clevis", "eye rod")),
    ("Mounts / Mounting Hardware", ("mount", "trunnion", "flange", "foot bracket", "pivot")),
    ("Bearings / Bushings", ("bearing", "bushing")),
    ("Jam Nuts", ("jam nut", "jamnut")),
    ("Pistons", ("piston",)),
    ("Barrels", ("barrel", "tube")),
    ("Fasteners", ("bolt", "screw", "washer", "stud", "fastener")),
    ("Rods", ("rod",)),
    ("Seals", ("seal", "o-ring", "oring", "packing", "wear ring", "backup ring")),
    ("Accessories", ("sensor", "cushion", "switch", "bracket", "accessory")),
]

SEAL_MATERIAL_GROUPS = [
    ("Viton", ("viton", "fkm")),
    ("Buna / Nitrile", ("buna", "nitrile", "nbr")),
    ("HNBR", ("hnbr", "highly saturated nitrile")),
    ("Polyurethane", ("polyurethane", "urethane")),
    ("Low Friction", ("low friction",)),
]

_SIZE_RE = re.compile(r"\d+(?:\.\d+)?")


def _part_text(row) -> str:
    return f"{row.description or ''} {row.all_descriptions or ''}".lower()


def categorize_part(row) -> str:
    """Return the display category for a part, preferring the stored value."""
    if getattr(row, "category", None):
        return row.category.strip()
    text = _part_text(row)
    for name, terms in _CATEGORY_KEYWORDS:
        if any(term in text for term in terms):
            return name
    return "Other / Unclassified"


def seal_material_group(row) -> str:
    text = _part_text(row)
    for name, terms in SEAL_MATERIAL_GROUPS:
        if any(term in text for term in terms):
            return name
    return "Other / Unclassified"


def part_sort_key(row):
    """Smallest-to-largest by any detected size; parts with no size sort last."""
    decoded = getattr(row, "family_size", None)
    if decoded is not None:
        return (0, float(decoded), row.part_number or "")
    for text in (row.part_number or "", row.description or ""):
        sizes = [float(match) for match in _SIZE_RE.findall(text) if float(match) > 0]
        if sizes:
            return (0, min(sizes), row.part_number or "")
    return (1, 0.0, row.part_number or "")


UNCLASSIFIED = "Other / Unclassified"
# Trailing bore/thread text such as " 1 1/2-12" or " 2 1/2" is the size, not the type.
_SIZE_SUFFIX_RE = re.compile(r"\s+[\d][\d/\-.\s]*$")


def part_subgroup(row) -> str:
    """Subgroup label taken from the part's own description, minus trailing size text."""
    text = (getattr(row, "description", None) or "").strip()
    label = _SIZE_SUFFIX_RE.sub("", text).strip()
    if not label:
        return UNCLASSIFIED
    last = label.rsplit(" ", 1)[-1]
    if last.isalpha() and len(last) > 2 and not last.isupper() and not label.endswith("s"):
        label = f"{label}s"
    return label


def _slugify(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-") or "category"


def _annotate_family_display(parts: list, families_by_code: dict) -> None:
    """Attach display-only family info to each part (nothing is persisted)."""
    for row in parts:
        family = families_by_code.get((getattr(row, "family_code", None) or "").upper())
        info = describe_part(family, row) if family else None
        row.family_display = family
        row.family_size = info["size"] if info else None
        row.size_display = info["size_display"] if info else None
        row.size_unresolved = bool(info["size_unresolved"]) if info else False
        row.generated_description = info["generated_description"] if info else None
        row.display_description = info["display_description"] if info else (row.description or "")


def _split_family_blocks(rows: list, families_by_code: dict) -> tuple[list, list]:
    """Split a group's rows into per-family blocks plus unassigned rows."""
    blocks: dict[str, list] = defaultdict(list)
    unassigned = []
    for row in rows:
        code = (getattr(row, "family_code", None) or "").upper()
        if code and code in families_by_code:
            blocks[code].append(row)
        else:
            unassigned.append(row)
    ordered = sorted(blocks, key=lambda code: families_by_code[code]["display_name"].lower())
    families = [
        {
            "family": families_by_code[code],
            "headline": family_headline(families_by_code[code]),
            "rows": sorted(blocks[code], key=part_sort_key),
        }
        for code in ordered
    ]
    return families, sorted(unassigned, key=part_sort_key)


def build_category_sections(parts: list, families_by_code: dict | None = None) -> list[dict]:
    """Group catalog parts into traditional part-family sections for the index."""
    families_by_code = families_by_code or {}
    _annotate_family_display(parts, families_by_code)
    by_category: dict[str, list] = defaultdict(list)
    for row in parts:
        by_category[categorize_part(row)].append(row)

    sections = []
    seen = set()
    ordered_names = CATEGORY_ORDER + [name for name in by_category if name not in CATEGORY_ORDER]
    for name in ordered_names:
        if name in seen or name not in by_category:
            continue
        seen.add(name)
        rows = by_category[name]
        sub: dict[str, list] = defaultdict(list)
        for row in rows:
            sub[part_subgroup(row)].append(row)
        sub_order = sorted((key for key in sub if key != UNCLASSIFIED), key=str.lower)
        if UNCLASSIFIED in sub:
            sub_order.append(UNCLASSIFIED)
        group_rows = [(group_name, sub[group_name]) for group_name in sub_order]
        groups = []
        for group_name, members in group_rows:
            families, unassigned = _split_family_blocks(members, families_by_code)
            groups.append({"name": group_name, "families": families, "rows": unassigned})
        sections.append({
            "name": name,
            "slug": _slugify(name),
            "count": len(rows),
            "groups": groups,
        })
    return sections


def catalog_page_data() -> dict[str, list]:
    """Return catalog rows for the admin page without touching pricing-engine data."""
    families = list_families()
    families_by_code = {family["family_code"].upper(): family for family in families}
    with get_session() as session:
        prices = session.execute(
            select(BaseAssemblyPrice).order_by(
                BaseAssemblyPrice.pricing_type,
                BaseAssemblyPrice.series,
                BaseAssemblyPrice.bore,
            )
        ).scalars().all()
        parts = session.execute(
            select(CatalogPart).order_by(CatalogPart.part_number)
        ).scalars().all()
        return {
            "parts": parts,
            "families": families,
            "size_rules": SIZE_RULES,
            "template_variables": TEMPLATE_VARIABLES,
            "category_sections": build_category_sections(parts, families_by_code),
            "base_prices": [row for row in prices if "assembly" not in row.pricing_type.lower()],
            "assembly_prices": [row for row in prices if "assembly" in row.pricing_type.lower()],
            "modifications": session.execute(
                select(CommonModificationPrice).order_by(
                    CommonModificationPrice.series,
                    CommonModificationPrice.option_name,
                )
            ).scalars().all(),
            "ph_va_prices": session.execute(
                select(PhVaPrice).order_by(
                    PhVaPrice.pricing_group,
                    PhVaPrice.series,
                    PhVaPrice.option_name,
                )
            ).scalars().all(),
            "format_money": format_money,
            "format_measurement": format_measurement,
            "format_price_unit": format_price_unit,
        }


def update_pricing_amount(
    table_name: str,
    row_id,
    field_name: str,
    value,
    actor=None,
) -> dict[str, str]:
    model, _ = _resolve_editable(table_name, field_name)
    parsed_row_id = _row_id(row_id)
    amount = _parse_amount(value)

    stored_amount = amount.quantize(STORAGE_QUANT)
    with get_session() as session:
        row = session.get(model, parsed_row_id)
        if row is None:
            raise PriceUpdateError("Pricing row was not found.", status_code=404)
        old_value = getattr(row, field_name)
        _apply_value(row, field_name, stored_amount)
        unit = getattr(row, "unit", None)
        _log_change(
            session,
            table_name=table_name,
            row=row,
            field_name=field_name,
            old_value=old_value,
            new_value=stored_amount,
            change_type="individual",
            actor=actor,
        )
        session.commit()

    return {
        "value": f"{stored_amount.quantize(MONEY_QUANT):.2f}",
        "formatted": f"{format_money(stored_amount)}{format_price_unit(unit) if field_name == 'per_unit_price' else ''}",
    }


# --- Bulk price updates -------------------------------------------------


BULK_METHODS = {
    "increase_percent",
    "decrease_percent",
    "add_amount",
    "subtract_amount",
    "set_exact",
}

METHOD_LABELS = {
    "increase_percent": "increased {amount}%",
    "decrease_percent": "decreased {amount}%",
    "add_amount": "increased by ${amount}",
    "subtract_amount": "decreased by ${amount}",
    "set_exact": "set to ${amount}",
}


def _resolve_editable(table_name: str, field_name: str):
    table_config = EDITABLE_PRICE_FIELDS.get((table_name or "").strip())
    if not table_config:
        raise PriceUpdateError("Unknown pricing table.")
    model, allowed_fields = table_config
    if field_name not in allowed_fields:
        raise PriceUpdateError("This pricing field is not editable.")
    return model, allowed_fields


def _row_id(row_id) -> int:
    try:
        return int(row_id)
    except (TypeError, ValueError) as exc:
        raise PriceUpdateError("Invalid pricing row.") from exc


def _parse_amount(value) -> Decimal:
    try:
        amount = Decimal(str(value).replace(",", "").replace("$", "").strip())
    except (InvalidOperation, AttributeError) as exc:
        raise PriceUpdateError("Enter a numeric price.") from exc
    if amount < 0:
        raise PriceUpdateError("Price must be non-negative.")
    return amount


def _apply_value(row, field_name: str, stored_amount: Decimal) -> None:
    setattr(row, field_name, stored_amount)
    raw_field = f"{field_name}_raw"
    if hasattr(row, raw_field):
        setattr(row, raw_field, None)


def _row_identity(table_name: str, row) -> tuple[str, str | None]:
    if table_name == "catalog_parts":
        return row.part_number, row.description
    label = getattr(row, "option_name", None) or getattr(row, "pricing_type", None) or f"#{row.id}"
    parts = [str(p) for p in (getattr(row, "series", None), getattr(row, "bore", None), getattr(row, "rod", None)) if p]
    return str(label), " / ".join(parts) or None


def _log_change(
    session: Session,
    *,
    table_name: str,
    row,
    field_name: str,
    old_value=None,
    new_value=None,
    change_type: str,
    actor=None,
    batch_id: str | None = None,
    batch_summary: str | None = None,
    old_text: str | None = None,
    new_text: str | None = None,
) -> None:
    label, description = _row_identity(table_name, row)
    session.add(
        PriceChangeLog(
            table_name=table_name,
            row_id=row.id,
            field_name=field_name,
            record_label=label,
            record_description=description,
            old_value=old_value,
            new_value=new_value,
            old_text=old_text,
            new_text=new_text,
            change_type=change_type,
            batch_id=batch_id,
            batch_summary=batch_summary,
            changed_by_user_id=getattr(actor, "id", None),
            changed_by_name=getattr(actor, "display_name", None),
        )
    )


def compute_new_price(current, method: str, amount: Decimal) -> Decimal:
    base = Decimal(str(current)) if current is not None else Decimal("0")
    if method == "set_exact":
        result = amount
    elif method == "increase_percent":
        result = base * (Decimal("1") + amount / Decimal("100"))
    elif method == "decrease_percent":
        result = base * (Decimal("1") - amount / Decimal("100"))
    elif method == "add_amount":
        result = base + amount
    else:
        result = base - amount
    if result < 0:
        result = Decimal("0")
    return result.quantize(STORAGE_QUANT)


def _validate_bulk(table_name: str, field_name: str, row_ids, method: str, amount) -> Decimal:
    _resolve_editable(table_name, field_name)
    if method not in BULK_METHODS:
        raise PriceUpdateError("Unknown bulk update method.")
    if not row_ids:
        raise PriceUpdateError("Select at least one record.")
    return _parse_amount(amount)


def preview_bulk_price_update(table_name: str, field_name: str, row_ids, method: str, amount) -> dict:
    parsed_amount = _validate_bulk(table_name, field_name, row_ids, method, amount)
    model, _ = _resolve_editable(table_name, field_name)
    ids = [_row_id(value) for value in row_ids]

    rows = []
    with get_session() as session:
        for row in session.execute(select(model).where(model.id.in_(ids))).scalars().all():
            current = getattr(row, field_name)
            if current is None:
                continue
            label, description = _row_identity(table_name, row)
            new_value = compute_new_price(current, method, parsed_amount)
            rows.append({
                "id": row.id,
                "label": label,
                "description": description or "",
                "old_value": f"{Decimal(str(current)).quantize(MONEY_QUANT):.2f}",
                "old_formatted": format_money(current),
                "new_value": f"{new_value.quantize(MONEY_QUANT):.2f}",
                "new_formatted": format_money(new_value),
                "conflicted": getattr(row, "review_needed", None) == "YES",
            })
    rows.sort(key=lambda item: item["label"])
    return {
        "rows": rows,
        "count": len(rows),
        "summary": describe_bulk_change(table_name, field_name, method, parsed_amount, len(rows)),
    }


def describe_bulk_change(table_name: str, field_name: str, method: str, amount: Decimal, count: int, context: str = "") -> str:
    amount_text = f"{amount.normalize():f}" if method.endswith("percent") else f"{amount.quantize(MONEY_QUANT):.2f}"
    field_text = field_name.replace("_", " ")
    scope = f"{context.strip()} " if context and context.strip() else ""
    return f"{scope}{field_text} {METHOD_LABELS[method].format(amount=amount_text)} ({count} records)".strip()


def apply_bulk_price_update(
    table_name: str,
    field_name: str,
    row_ids,
    method: str,
    amount,
    actor=None,
    context: str = "",
) -> dict:
    parsed_amount = _validate_bulk(table_name, field_name, row_ids, method, amount)
    model, _ = _resolve_editable(table_name, field_name)
    ids = [_row_id(value) for value in row_ids]
    batch_id = uuid.uuid4().hex

    with get_session() as session:
        rows = session.execute(select(model).where(model.id.in_(ids))).scalars().all()
        targets = [row for row in rows if getattr(row, field_name) is not None]
        if not targets:
            raise PriceUpdateError("No selected record has a price to update.")
        summary = describe_bulk_change(table_name, field_name, method, parsed_amount, len(targets), context)
        updated = []
        for row in targets:
            old_value = getattr(row, field_name)
            new_value = compute_new_price(old_value, method, parsed_amount)
            _apply_value(row, field_name, new_value)
            _log_change(
                session,
                table_name=table_name,
                row=row,
                field_name=field_name,
                old_value=old_value,
                new_value=new_value,
                change_type="bulk",
                actor=actor,
                batch_id=batch_id,
                batch_summary=summary,
            )
            updated.append({
                "id": row.id,
                "field": field_name,
                "value": f"{new_value.quantize(MONEY_QUANT):.2f}",
                "formatted": format_money(new_value),
            })
        session.commit()

    return {"updated": updated, "count": len(updated), "batch_id": batch_id, "summary": summary}


# --- History / restore --------------------------------------------------


def recent_price_changes(limit: int = 50) -> list[dict]:
    with get_session() as session:
        entries = session.execute(
            select(PriceChangeLog).order_by(PriceChangeLog.id.desc()).limit(max(1, min(int(limit), 500)))
        ).scalars().all()
        return [
            {
                "id": entry.id,
                "table": entry.table_name,
                "row_id": entry.row_id,
                "field": entry.field_name,
                "field_label": entry.field_name.replace("_", " ").title(),
                "label": entry.record_label or "",
                "description": entry.record_description or "",
                "old_formatted": format_money(entry.old_value) if entry.old_value is not None else (entry.old_text or ""),
                "new_formatted": format_money(entry.new_value) if entry.new_value is not None else (entry.new_text or ""),
                "change_type": entry.change_type,
                "batch_summary": entry.batch_summary or "",
                "changed_by": entry.changed_by_name or "Unknown",
                "created_at": entry.created_at.strftime("%Y-%m-%d %H:%M"),
                "restorable": entry.old_value is not None,
            }
            for entry in entries
        ]


def restore_price_change(change_id, actor=None) -> dict:
    try:
        parsed_id = int(change_id)
    except (TypeError, ValueError) as exc:
        raise PriceUpdateError("Invalid history entry.") from exc

    with get_session() as session:
        entry = session.get(PriceChangeLog, parsed_id)
        if entry is None:
            raise PriceUpdateError("History entry was not found.", status_code=404)
        if entry.old_value is None:
            raise PriceUpdateError("This change has no previous price to restore.")
        model, _ = _resolve_editable(entry.table_name, entry.field_name)
        row = session.get(model, entry.row_id)
        if row is None:
            raise PriceUpdateError("Pricing row was not found.", status_code=404)

        current = getattr(row, entry.field_name)
        restored = Decimal(str(entry.old_value)).quantize(STORAGE_QUANT)
        _apply_value(row, entry.field_name, restored)
        _log_change(
            session,
            table_name=entry.table_name,
            row=row,
            field_name=entry.field_name,
            old_value=current,
            new_value=restored,
            change_type="restore",
            actor=actor,
            batch_id=entry.batch_id,
            batch_summary=f"Restored change #{entry.id}",
        )
        session.commit()

    return {
        "table": entry.table_name,
        "id": entry.row_id,
        "field": entry.field_name,
        "value": f"{restored.quantize(MONEY_QUANT):.2f}",
        "formatted": format_money(restored),
    }


# --- Conflict resolution --------------------------------------------------


# Maps an editable field to the (label, source-column) pairs that may conflict.
CONFLICT_SOURCE_FIELDS = {
    "sell_price": (("ShipVia", "shipvia_sell_values"), ("Acc", "acc_sell_values")),
    "unit_cost": (("Inventory", "inventory_cost_values"),),
}


def get_conflict_details(row_id, field: str) -> dict:
    if field not in CONFLICT_SOURCE_FIELDS:
        raise PriceUpdateError("This field does not support conflict resolution.")
    with get_session() as session:
        row = session.get(CatalogPart, _row_id(row_id))
        if row is None:
            raise PriceUpdateError("Part was not found.", status_code=404)
        sources = []
        for label, attr in CONFLICT_SOURCE_FIELDS[field]:
            value = _decimal(getattr(row, attr, None))
            if value is None:
                continue
            sources.append({"name": label, "value": f"{value:.2f}", "formatted": format_money(value)})
        current = getattr(row, field)
        return {
            "part_number": row.part_number,
            "description": row.description or "",
            "field": field,
            "field_label": field.replace("_", " ").title(),
            "current_formatted": format_money(current),
            "sources": sources,
            "has_conflict": row.review_needed == "YES",
            "review_reason": row.review_reason or "",
        }


def resolve_conflict(row_id, field: str, resolution: str, source_name, value, actor=None) -> dict:
    if field not in CONFLICT_SOURCE_FIELDS:
        raise PriceUpdateError("This field does not support conflict resolution.")
    if resolution not in {"use_source", "custom", "keep"}:
        raise PriceUpdateError("Choose a resolution option.")

    with get_session() as session:
        row = session.get(CatalogPart, _row_id(row_id))
        if row is None:
            raise PriceUpdateError("Part was not found.", status_code=404)

        if resolution == "keep":
            current = getattr(row, field)
            return {
                "table": "catalog_parts",
                "id": row.id,
                "field": field,
                "value": f"{Decimal(str(current)).quantize(MONEY_QUANT):.2f}" if current is not None else "",
                "formatted": format_money(current),
                "resolved": False,
            }

        if resolution == "use_source":
            source_attrs = dict(CONFLICT_SOURCE_FIELDS[field])
            attr = source_attrs.get(source_name)
            if not attr:
                raise PriceUpdateError("Unknown source price.")
            new_value = _decimal(getattr(row, attr, None))
            if new_value is None:
                raise PriceUpdateError("That source has no stored price.")
            resolution_note = f"used {source_name} price"
        else:
            new_value = _parse_amount(value)
            resolution_note = "used a custom price"

        old_value = getattr(row, field)
        stored_amount = new_value.quantize(STORAGE_QUANT)
        _apply_value(row, field, stored_amount)
        row.review_needed = "NO"
        actor_name = getattr(actor, "display_name", None) or "Unknown"
        stamp = utc_now().strftime("%Y-%m-%d %H:%M")
        note = f" -- Resolved by {actor_name} on {stamp}: {resolution_note}"
        row.review_reason = f"{(row.review_reason or '').strip()}{note}".strip()
        _log_change(
            session,
            table_name="catalog_parts",
            row=row,
            field_name=field,
            old_value=old_value,
            new_value=stored_amount,
            change_type="conflict_resolved",
            actor=actor,
            batch_summary=f"Conflict resolved: {resolution_note}",
        )
        session.commit()

    return {
        "table": "catalog_parts",
        "id": row.id,
        "field": field,
        "value": f"{stored_amount.quantize(MONEY_QUANT):.2f}",
        "formatted": format_money(stored_amount),
        "resolved": True,
    }


# --- Add Part --------------------------------------------------------------


def _validate_optional_amount(value, label: str) -> Decimal | None:
    if value in (None, ""):
        return None
    try:
        amount = Decimal(str(value).replace(",", "").replace("$", "").strip())
    except InvalidOperation as exc:
        raise PriceUpdateError(f"{label} must be a number.") from exc
    if amount < 0:
        raise PriceUpdateError(f"{label} must be non-negative.")
    return amount.quantize(STORAGE_QUANT)


def create_catalog_part(data: dict, actor=None) -> dict:
    part_number = _clean(data.get("part_number"))
    if not part_number:
        raise PriceUpdateError("Part Number is required.")
    description = _clean(data.get("description"))
    if not description:
        raise PriceUpdateError("Description is required.")
    category = _clean(data.get("category"))
    if not category:
        raise PriceUpdateError("Category is required.")
    unit_cost = _validate_optional_amount(data.get("unit_cost"), "Unit Cost")
    sell_price = _validate_optional_amount(data.get("sell_price"), "Sell Price")
    sell_price_source = _clean(data.get("sell_price_source")) or "Manual"
    vendors = _normalize_vendors(data.get("vendors"))
    source_locations = _clean(data.get("source_locations"))
    active = "YES" if data.get("active", True) in (True, "true", "YES", "yes", "1", 1) else "NO"

    with get_session() as session:
        existing = session.execute(
            select(CatalogPart).filter_by(part_number=part_number)
        ).scalar_one_or_none()
        if existing is not None:
            raise PriceUpdateError(f"Part number '{part_number}' already exists.")

        row = CatalogPart(
            part_number=part_number,
            description=description,
            category=category,
            unit_cost=unit_cost,
            sell_price=sell_price,
            sell_price_source=sell_price_source,
            vendors=vendors,
            source_locations=source_locations,
            active=active,
            review_needed="NO",
        )
        session.add(row)
        try:
            session.flush()
        except IntegrityError as exc:
            session.rollback()
            raise PriceUpdateError(f"Part number '{part_number}' already exists.") from exc

        _log_change(
            session,
            table_name="catalog_parts",
            row=row,
            field_name="sell_price",
            old_value=None,
            new_value=sell_price,
            change_type="created",
            actor=actor,
            batch_summary=f"Created part {part_number}",
        )
        session.commit()
        return {"id": row.id, "part_number": row.part_number, "category": row.category}


# --- Edit an existing part -------------------------------------------------


# Text fields an admin may edit inline, mapped to their human label.
EDITABLE_PART_TEXT_FIELDS = {
    "part_number": "Part Number",
    "description": "Description",
    "category": "Category",
    "vendors": "Vendors",
    "source_locations": "Workbook Location",
    "active": "Active",
}

EDITABLE_PART_AMOUNT_FIELDS = {"unit_cost": "Unit Cost", "sell_price": "Sell Price"}


def _validate_inventory(value) -> int | None:
    if value in (None, ""):
        return None
    try:
        inventory = int(str(value).strip())
    except (TypeError, ValueError) as exc:
        raise PriceUpdateError("Inventory must be a whole number.") from exc
    if inventory < 0:
        raise PriceUpdateError("Inventory must be non-negative.")
    return inventory


def _normalize_vendors(value) -> str | None:
    """Store a multi-vendor list as one vendor per line."""
    if value is None:
        return None
    text = str(value).replace("\r\n", "\n").replace("\r", "\n")
    entries = [part.strip() for chunk in text.split("\n") for part in chunk.split(";")]
    return "\n".join(entry for entry in entries if entry) or None


def search_catalog_parts(query: str, limit: int = 25) -> list[dict]:
    """Lightweight lookup for non-admin pages (e.g. Order Form "Add Part")."""
    query = _clean(query) or ""
    with get_session() as session:
        statement = select(CatalogPart).order_by(CatalogPart.part_number)
        if query:
            like = f"%{query}%"
            statement = statement.where(
                CatalogPart.part_number.ilike(like) | CatalogPart.description.ilike(like)
            )
        rows = session.execute(statement.limit(max(1, min(limit, 100)))).scalars().all()
        return [
            {
                "id": row.id,
                "part_number": row.part_number,
                "description": row.description or "",
                "unit_cost": f"{Decimal(str(row.unit_cost)):.2f}" if row.unit_cost is not None else "",
                "sell_price": f"{Decimal(str(row.sell_price)):.2f}" if row.sell_price is not None else "",
            }
            for row in rows
        ]


def quote_special_parts(allowed_part_numbers) -> list[dict]:
    """Return active catalog details for special parts known to the pricing engine."""
    allowed = {str(value).strip().upper() for value in allowed_part_numbers if str(value).strip()}
    if not allowed:
        return []
    with get_session() as session:
        rows = session.execute(select(CatalogPart).order_by(CatalogPart.part_number)).scalars().all()
        return [
            {
                "part_number": row.part_number,
                "description": row.description or "",
                "sell_price": f"{Decimal(str(row.sell_price)):.2f}" if row.sell_price is not None else "",
            }
            for row in rows
            if (row.active or "YES").upper() != "NO"
            and (row.part_number or "").strip().upper() in allowed
        ]


def get_catalog_part(row_id) -> dict:
    with get_session() as session:
        row = session.get(CatalogPart, _row_id(row_id))
        if row is None:
            raise PriceUpdateError("Part was not found.", status_code=404)
        return {
            "id": row.id,
            "part_number": row.part_number,
            "description": row.description or "",
            "category": categorize_part(row),
            "vendors": row.vendors or "",
            "source_locations": row.source_locations or "",
            "active": (row.active or "YES").upper() != "NO",
            "unit_cost": f"{Decimal(str(row.unit_cost)):.2f}" if row.unit_cost is not None else "",
            "sell_price": f"{Decimal(str(row.sell_price)):.2f}" if row.sell_price is not None else "",
            "inventory": row.inventory if row.inventory is not None else "",
        }


def update_catalog_part(row_id, data: dict, actor=None) -> dict:
    """Edit any admin-editable field on one catalog part, logging each change."""
    with get_session() as session:
        row = session.get(CatalogPart, _row_id(row_id))
        if row is None:
            raise PriceUpdateError("Part was not found.", status_code=404)

        changes: list[tuple[str, str | None, str | None]] = []
        for field, label in EDITABLE_PART_TEXT_FIELDS.items():
            if field not in data:
                continue
            if field == "active":
                new_text = "YES" if data.get("active") in (True, "true", "YES", "yes", "1", 1) else "NO"
            elif field == "vendors":
                new_text = _normalize_vendors(data.get(field))
            else:
                new_text = _clean(data.get(field))
            if field in {"part_number", "description", "category"} and not new_text:
                raise PriceUpdateError(f"{label} is required.")
            old_text = getattr(row, field)
            if (old_text or None) == (new_text or None):
                continue
            if field == "part_number":
                clash = session.execute(
                    select(CatalogPart).filter_by(part_number=new_text)
                ).scalar_one_or_none()
                if clash is not None and clash.id != row.id:
                    raise PriceUpdateError(f"Part number '{new_text}' already exists.")
            setattr(row, field, new_text)
            changes.append((field, old_text, new_text))

        amount_changes: list[tuple[str, Decimal | None, Decimal | None]] = []
        for field, label in EDITABLE_PART_AMOUNT_FIELDS.items():
            if field not in data:
                continue
            new_amount = _validate_optional_amount(data.get(field), label)
            old_amount = getattr(row, field)
            if old_amount is None and new_amount is None:
                continue
            if old_amount is not None and new_amount is not None and Decimal(str(old_amount)) == new_amount:
                continue
            setattr(row, field, new_amount)
            amount_changes.append((field, old_amount, new_amount))

        if "inventory" in data:
            new_inventory = _validate_inventory(data.get("inventory"))
            old_inventory = row.inventory
            if old_inventory != new_inventory:
                row.inventory = new_inventory
                amount_changes.append(("inventory", old_inventory, new_inventory))

        if not changes and not amount_changes:
            return {"id": row.id, "changed": 0, **_part_row_display(row)}

        for field, old_text, new_text in changes:
            _log_change(
                session,
                table_name="catalog_parts",
                row=row,
                field_name=field,
                change_type="edit",
                actor=actor,
                old_text=old_text,
                new_text=new_text,
                batch_summary=f"Edited {EDITABLE_PART_TEXT_FIELDS[field]}",
            )
        for field, old_amount, new_amount in amount_changes:
            _log_change(
                session,
                table_name="catalog_parts",
                row=row,
                field_name=field,
                old_value=old_amount,
                new_value=new_amount,
                change_type="edit",
                actor=actor,
                batch_summary=f"Edited {EDITABLE_PART_AMOUNT_FIELDS.get(field, 'Inventory')}",
            )
        try:
            session.commit()
        except IntegrityError as exc:
            session.rollback()
            raise PriceUpdateError("That part number is already in use.") from exc

        return {"id": row.id, "changed": len(changes) + len(amount_changes), **_part_row_display(row)}


def _part_row_display(row) -> dict:
    return {
        "part_number": row.part_number,
        "description": row.description or "",
        "vendors": row.vendors or "",
        "source_locations": row.source_locations or "",
        "unit_cost": f"{Decimal(str(row.unit_cost)).quantize(MONEY_QUANT):.2f}" if row.unit_cost is not None else "",
        "unit_cost_formatted": format_money(row.unit_cost),
        "sell_price": f"{Decimal(str(row.sell_price)).quantize(MONEY_QUANT):.2f}" if row.sell_price is not None else "",
        "sell_price_formatted": format_money(row.sell_price),
        "inventory": row.inventory if row.inventory is not None else "",
    }


def update_catalog_inventory_bulk(row_ids, value, actor=None) -> dict:
    inventory = _validate_inventory(value)
    ids = [_row_id(row_id) for row_id in (row_ids or [])]
    if not ids:
        raise PriceUpdateError("Select at least one part.")
    with get_session() as session:
        rows = session.execute(select(CatalogPart).where(CatalogPart.id.in_(ids))).scalars().all()
        if len(rows) != len(set(ids)):
            raise PriceUpdateError("One or more selected parts were not found.", status_code=404)
        updated = []
        for row in rows:
            if row.inventory == inventory:
                continue
            old_inventory = row.inventory
            row.inventory = inventory
            _log_change(
                session,
                table_name="catalog_parts",
                row=row,
                field_name="inventory",
                old_value=old_inventory,
                new_value=inventory,
                change_type="edit",
                actor=actor,
                batch_summary="Bulk inventory update",
            )
            updated.append({"id": row.id, "inventory": inventory})
        session.commit()
        return {"updated": updated, "count": len(updated)}


# --- Excel price workbook upload -------------------------------------------


WORKBOOK_FIELD_LABELS = {"unit_cost": "Unit Cost", "sell_price": "Sell Price", "both": "Unit Cost and Sell Price"}
WORKBOOK_PRICE_FIELDS = ("unit_cost", "sell_price")

_PART_NUMBER_HEADERS = ("part number", "part no", "part #", "part num", "partnumber", "part", "item number", "item")
_DESCRIPTION_HEADERS = ("part name", "description", "name", "desc")
_PRICE_HEADERS = {
    "sell_price": ("sell price", "new price", "price", "list price", "sell", "unit price"),
    "unit_cost": ("unit cost", "new cost", "cost", "our cost", "purchase price"),
}
_MAX_WORKBOOK_ROWS = 20000


def _header_text(value) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip().lower())


def _find_column(headers: list[str], candidates: tuple[str, ...]) -> int | None:
    for candidate in candidates:
        for index, header in enumerate(headers):
            if header == candidate:
                return index
    for candidate in candidates:
        for index, header in enumerate(headers):
            if candidate in header:
                return index
    return None


def _cell_text(value) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return _clean(str(value))


def _workbook_price(value) -> Decimal | None:
    if value in (None, ""):
        return None
    text = str(value).replace(",", "").replace("$", "").strip()
    if not text:
        return None
    try:
        amount = Decimal(text)
    except InvalidOperation:
        return None
    return amount if amount >= 0 else None


def parse_price_workbook(stream, field: str) -> dict:
    """Read an uploaded .xlsx price list and diff it against the catalog."""
    from openpyxl import load_workbook  # imported lazily: only needed for uploads

    if field not in WORKBOOK_FIELD_LABELS:
        raise PriceUpdateError("Choose Unit Cost or Sell Price before uploading.")
    if stream is None:
        raise PriceUpdateError("Choose an Excel (.xlsx) file to upload.")

    try:
        workbook = load_workbook(stream, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many types for bad input
        raise PriceUpdateError("That file could not be read as an Excel (.xlsx) workbook.") from exc

    try:
        sheet = workbook.worksheets[0]
        rows = []
        for index, raw in enumerate(sheet.iter_rows(values_only=True)):
            rows.append(raw)
            if index >= _MAX_WORKBOOK_ROWS:
                break
    finally:
        workbook.close()

    header_index = None
    part_col = desc_col = None
    price_cols = {}
    for index, raw in enumerate(rows[:10]):
        headers = [_header_text(cell) for cell in raw]
        part_col = _find_column(headers, _PART_NUMBER_HEADERS)
        desc_col = _find_column(headers, _DESCRIPTION_HEADERS)
        fields = WORKBOOK_PRICE_FIELDS if field == "both" else (field,)
        price_cols = {name: _find_column(headers, _PRICE_HEADERS[name]) for name in fields}
        if any(column is not None for column in price_cols.values()) and (part_col is not None or desc_col is not None):
            header_index = index
            break
    if header_index is None:
        raise PriceUpdateError(
            "The workbook needs a header row with a part number (or part name) column "
            f"and a {WORKBOOK_FIELD_LABELS[field]} column."
        )

    entries = []
    for raw in rows[header_index + 1:]:
        key = _cell_text(raw[part_col]) if part_col is not None and part_col < len(raw) else None
        name = _cell_text(raw[desc_col]) if desc_col is not None and desc_col < len(raw) else None
        prices = {
            name: _workbook_price(raw[column]) if column is not None and column < len(raw) else None
            for name, column in price_cols.items()
        }
        if not key and not name:
            continue
        entries.append((key, name, prices))

    with get_session() as session:
        parts = session.execute(select(CatalogPart)).scalars().all()
        by_number = {(part.part_number or "").strip().lower(): part for part in parts}
        by_description = {}
        for part in parts:
            description_key = (part.description or "").strip().lower()
            if description_key:
                by_description.setdefault(description_key, part)

        matched, unchanged, invalid = [], 0, []
        seen_ids = set()
        for key, name, prices in entries:
            part = by_number.get((key or "").lower()) or by_description.get((name or "").lower())
            if part is None:
                invalid.append({"reason": "not found", "value": key or name})
                continue
            valid_prices = {name: price for name, price in prices.items() if price is not None}
            if not valid_prices:
                invalid.append({"reason": "no valid price", "value": key or name})
                continue
            if part.id in seen_ids:
                continue
            seen_ids.add(part.id)
            changes = {}
            for price_field, price in valid_prices.items():
                current = getattr(part, price_field)
                new_value = price.quantize(STORAGE_QUANT)
                if current is None or Decimal(str(current)).quantize(STORAGE_QUANT) != new_value:
                    changes[price_field] = {
                        "old_value": f"{Decimal(str(current)).quantize(MONEY_QUANT):.2f}" if current is not None else "",
                        "old_formatted": format_money(current) or "(none)",
                        "new_value": f"{new_value.quantize(MONEY_QUANT):.2f}",
                        "new_formatted": format_money(new_value),
                    }
            if not changes:
                unchanged += 1
                continue
            if field == "both":
                old_display = " / ".join(f"{WORKBOOK_FIELD_LABELS[key]}: {value['old_formatted']}" for key, value in changes.items())
                new_display = " / ".join(f"{WORKBOOK_FIELD_LABELS[key]}: {value['new_formatted']}" for key, value in changes.items())
                matched.append({
                    "id": part.id,
                    "label": part.part_number,
                    "description": part.description or "",
                    "old_formatted": old_display,
                    "new_formatted": new_display,
                    "values": {key: value["new_value"] for key, value in changes.items()},
                    "conflicted": part.review_needed == "YES",
                })
                continue
            price_change = changes[field]
            matched.append({
                "id": part.id,
                "label": part.part_number,
                "description": part.description or "",
                "old_value": price_change["old_value"],
                "old_formatted": price_change["old_formatted"],
                "new_value": price_change["new_value"],
                "new_formatted": price_change["new_formatted"],
                "conflicted": part.review_needed == "YES",
            })

    matched.sort(key=lambda item: item["label"])
    return {
        "rows": matched,
        "count": len(matched),
        "unchanged": unchanged,
        "unmatched": len(invalid),
        "unmatched_samples": [item["value"] for item in invalid[:10] if item["value"]],
        "field": field,
    }


def apply_workbook_price_update(field: str, updates, actor=None, source_name: str = "") -> dict:
    """Apply the reviewed rows from an uploaded price workbook."""
    fields = WORKBOOK_PRICE_FIELDS if field == "both" else (field,)
    for price_field in fields:
        _resolve_editable("catalog_parts", price_field)
    parsed = []
    for item in updates or []:
        item = item or {}
        if field == "both":
            values = {
                price_field: _parse_amount(value)
                for price_field, value in (item.get("values") or {}).items()
                if price_field in WORKBOOK_PRICE_FIELDS
            }
            parsed.append((_row_id(item.get("id")), values))
        else:
            parsed.append((_row_id(item.get("id")), _parse_amount(item.get("value"))))
    if not parsed:
        raise PriceUpdateError("Select at least one part to update.")

    by_id = {row_id: amount for row_id, amount in parsed}
    batch_id = uuid.uuid4().hex
    label = _clean(source_name) or "uploaded workbook"
    summary = f"{WORKBOOK_FIELD_LABELS[field]} updated from {label} ({len(by_id)} records)"

    with get_session() as session:
        rows = session.execute(select(CatalogPart).where(CatalogPart.id.in_(list(by_id)))).scalars().all()
        if not rows:
            raise PriceUpdateError("None of the selected parts were found.", status_code=404)
        updated = []
        for row in rows:
            changes = by_id[row.id] if field == "both" else {field: by_id[row.id]}
            for price_field, amount in changes.items():
                if amount is None:
                    raise PriceUpdateError(f"{WORKBOOK_FIELD_LABELS[price_field]} must be a valid non-negative amount.")
                new_value = amount.quantize(STORAGE_QUANT)
                old_value = getattr(row, price_field)
                if old_value is not None and Decimal(str(old_value)).quantize(STORAGE_QUANT) == new_value:
                    continue
                _apply_value(row, price_field, new_value)
                _log_change(
                    session,
                    table_name="catalog_parts",
                    row=row,
                    field_name=price_field,
                    old_value=old_value,
                    new_value=new_value,
                    change_type="upload",
                    actor=actor,
                    batch_id=batch_id,
                    batch_summary=summary,
                )
                updated.append({
                    "id": row.id,
                    "field": price_field,
                    "value": f"{new_value.quantize(MONEY_QUANT):.2f}",
                    "formatted": format_money(new_value),
                })
        session.commit()

    return {"updated": updated, "count": len(updated), "batch_id": batch_id, "summary": summary}


PRICE_TEMPLATE_HEADERS = ["Part Number", "Part Name", "Unit Cost", "Sell Price", "Vendors"]


def build_price_template_workbook():
    """Return an in-memory .xlsx upload template matching the catalog columns."""
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Font

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Price Update"
    sheet.append(PRICE_TEMPLATE_HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for column, width in zip("ABCDE", (22, 44, 14, 14, 40)):
        sheet.column_dimensions[column].width = width

    with get_session() as session:
        parts = session.execute(
            select(CatalogPart).order_by(CatalogPart.part_number)
        ).scalars().all()
        for part in parts:
            sheet.append([
                part.part_number,
                part.description or "",
                float(part.unit_cost) if part.unit_cost is not None else None,
                float(part.sell_price) if part.sell_price is not None else None,
                (part.vendors or "").replace("\n", "; "),
            ])

    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    buffer.seek(0)
    return buffer

